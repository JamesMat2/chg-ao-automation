import anthropic
import json
import os
import re
from datetime import date
import math
import requests

# Session 25: raised from 50000 after AO 2026-55-109 (168979 chars extracted)
# was silently truncated mid-criteria. ~75k tokens, well within context window.
MAX_AO_CHARS = 300000


class TruncatedResponseError(Exception):
    """Session 52: raised when a Claude response's stop_reason is 'max_tokens'
    -- a confirmed truncation signal read directly from the API, not inferred
    after the fact from a downstream JSON-parse failure. Every heuristic that
    catches truncation today (JSON-parse retries, the criteres_detail
    completeness check) only recognizes specific symptoms of it; this checks
    the actual cause directly."""


class ResponseRefusedError(Exception):
    """Session 52: raised when a Claude response's stop_reason is 'refusal'
    (or content is otherwise empty) -- confirmed via a real production
    incident (AO 13203, "Centre national de primatologie pour la
    preparation aux pandemies") where the facility-name/topic framing
    triggered a deterministic, reproducible refusal (confirmed 3/3 live
    attempts). Distinct from TruncatedResponseError: no retry-with-more-
    tokens can fix a refusal -- the input itself needs to change (see the
    Session 52 name-substitution test) or the AO needs manual handling."""


# Session 52: explicit, hand-curated list of (real_phrase, placeholder) pairs
# for AO content confirmed to trigger a Claude refusal (stop_reason ==
# "refusal"). Not automatic detection -- reliably identifying WHICH phrase
# caused a refusal without asking the model itself is circular (that call
# could refuse too). Extend by hand as new real cases are confirmed via the
# same reproduction method used for AO 13203 (DVR-20090): download the
# principal document, replay the exact failing call against the real API a
# few times to confirm it's deterministic, then add the trigger phrase here.
# Placeholders are deliberately distinctive (bracketed, uppercase) so
# restoration can't be confused with genuine AO content.
REFUSAL_SUBSTITUTIONS = [
    ("Centre national de primatologie pour la préparation aux pandémies", "[AO-ORGANISME-A]"),
]


def apply_refusal_substitutions(text):
    """Replace known refusal-trigger phrases with their placeholder tag, for
    a one-shot retry after a confirmed Claude refusal. See
    REFUSAL_SUBSTITUTIONS."""
    for phrase, placeholder in REFUSAL_SUBSTITUTIONS:
        text = text.replace(phrase, placeholder)
    return text


def restore_refusal_substitutions(obj):
    """Recursively walk obj (str/dict/list) and replace placeholder tags
    back with their real phrase. Safe to call unconditionally -- a no-op if
    no substitution was ever applied, since the placeholder tags are
    distinctive enough not to occur in genuine AO content. Call exactly once
    per script, right before data that may have been built from substituted
    text leaves that script's boundary (returned, persisted, or written into
    a document) -- not after every individual Claude call, which would
    re-expose downstream calls in the same script to the same trigger."""
    if isinstance(obj, str):
        for phrase, placeholder in REFUSAL_SUBSTITUTIONS:
            obj = obj.replace(placeholder, phrase)
        return obj
    if isinstance(obj, dict):
        return {k: restore_refusal_substitutions(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [restore_refusal_substitutions(v) for v in obj]
    return obj

# ── Text extraction ────────────────────────────────────────────────────────────

def extract_text_pdf(path):
    try:
        from pdfminer.high_level import extract_text
        text = extract_text(path)
        if text and len(text) > MAX_AO_CHARS:
            print(f"  [extraction] ATTENTION: document PDF tronque de {len(text)} a {MAX_AO_CHARS} caracteres")
        return text[:MAX_AO_CHARS] if text else ""
    except Exception as e:
        print(f"Erreur extraction PDF: {e}")
        return ""

def extract_text_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        full_text = "\n".join(parts)
        if len(full_text) > MAX_AO_CHARS:
            print(f"  [extraction] ATTENTION: document DOCX tronque de {len(full_text)} a {MAX_AO_CHARS} caracteres")
        return full_text[:MAX_AO_CHARS]
    except Exception as e:
        print(f"Erreur extraction DOCX: {e}")
        return ""

def extract_text_excel(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"[Feuille: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
                if cells:
                    parts.append(" | ".join(cells))
        full_text = "\n".join(parts)
        if len(full_text) > MAX_AO_CHARS:
            print(f"  [extraction] ATTENTION: document Excel tronque de {len(full_text)} a {MAX_AO_CHARS} caracteres")
        return full_text[:MAX_AO_CHARS]
    except Exception as e:
        print(f"Erreur extraction Excel: {e}")
        return ""

def extract_ao_text(path):
    if not path:
        return ""
    ext = os.path.splitext(path)[1].lower()
    if ext in ['.docx', '.doc']:
        text = extract_text_docx(path)
    elif ext in ['.xlsx', '.xls']:
        text = extract_text_excel(path)
    else:
        text = extract_text_pdf(path)
    return text

# ── Post-processor: normalize criteres_detail keys ────────────────────────────

def normalize_criteres_detail(criteres):
    """
    Claude sometimes returns variant key names for criteres_detail.
    This normalizes all variants to the exact keys the renderer expects:
      numero_critere, intitule, points_max, contenu_requis,
      nb_projets_requis, annees_experience_requises

    Session 50 fix: Claude sometimes returns criteres_detail as a list of
    plain strings instead of objects — either a document checklist when no
    scored grid exists (invitation-based AOs), or a placeholder describing
    an evaluation grid that exists but wasn't in the extracted text window.
    Non-dict entries are wrapped as {"intitule": str(c)} before key
    extraction so the schema (and every downstream consumer) always gets
    list[dict], instead of this function raising AttributeError on c.get(...)
    and getting silently swallowed by the caller's try/except, leaving the
    malformed list untouched.

    Session 51 fix: guards the CONTAINER shape too, not just entries within
    it. A single dict (one criterion returned bare instead of in a
    one-element list) is wrapped as [criteres] -- preserving its real data,
    since discarding it (or iterating its keys, which silently produces
    junk entries named after the dict's own field names) would lose real
    extracted content. A non-empty string is wrapped as a one-element
    placeholder, same spirit as the entry-level fix above. Anything else
    unexpected defaults to [].
    """
    if isinstance(criteres, dict):
        criteres = [criteres]
    elif isinstance(criteres, str):
        criteres = [criteres] if criteres.strip() else []
    elif not isinstance(criteres, list):
        criteres = []

    normalized = []
    for c in criteres:
        if not isinstance(c, dict):
            c = {"intitule": str(c)}
        n = {}
        # numero_critere
        n["numero_critere"] = (
            c.get("numero_critere") or
            c.get("numero") or
            c.get("critere") or
            c.get("id") or ""
        )
        # intitule
        n["intitule"] = (
            c.get("intitule") or
            c.get("titre") or
            c.get("nom") or
            c.get("name") or ""
        )
        # points_max
        raw_pts = (
            c.get("points_max") or
            c.get("points") or
            c.get("ponderation") or
            c.get("weight") or 0
        )
        try:
            n["points_max"] = int(raw_pts)
        except (ValueError, TypeError):
            n["points_max"] = 0
        # contenu_requis
        n["contenu_requis"] = (
            c.get("contenu_requis") or
            c.get("description") or
            c.get("contenu") or
            c.get("details") or
            c.get("exigences") or ""
        )
        # nb_projets_requis
        try:
            n["nb_projets_requis"] = int(c.get("nb_projets_requis") or c.get("nb_projets") or 0)
        except (ValueError, TypeError):
            n["nb_projets_requis"] = 0
        # annees_experience_requises
        try:
            n["annees_experience_requises"] = int(
                c.get("annees_experience_requises") or
                c.get("annees_experience") or
                c.get("experience_years") or 0
            )
        except (ValueError, TypeError):
            n["annees_experience_requises"] = 0

        normalized.append(n)
    return normalized

# ── Post-processor: classify expertise against CHG's 12 service categories ──
# Session 50 (Column E reframe, replacing the bordereau-sourced approach
# entirely per Alexandra/Pat): a dedicated, small, focused call rather than
# folded into the main analyze_ao_gonogo() extraction, which already carries
# the full ~300k-char AO text and competes for its 8000-token output budget
# across 9 unrelated schema sections -- the same prompt-competition pattern
# that caused Column 9's silent defaults. This call needs only the
# already-extracted fields below, so input cost is trivial.

CHG_SERVICE_CATEGORIES = [
    "Infrastructure",
    "Structure, bâtiment et ouvrage d'art",
    "Hydrologie et hydraulique",
    "Arpentage, drone et bathymétrie",
    "Transport",
    "Gestion de projet",
    "Environnement",
    "Études économiques",
    "Foresterie",
    "Géomatique et technologie de l'information",
    "Planification stratégique",
    "Urbanisme et aménagement du territoire",
]


def validate_expertise_categories(raw_categories):
    """
    Enforce that every returned category is either one of CHG's 12 known
    service categories (exact match, case/whitespace-normalized -- output
    uses the canonical spelling) or an explicit "autre: X" entry. Never
    silently coerced into the nearest known category, never silently
    dropped -- same principle as dedup_expertises()'s cross-list-only
    restriction and the criteres_detail fix earlier this session: the
    failure mode to design against is silent drift/hallucination, not
    visible "autre" noise.
    """
    canonical_by_key = {c.strip().lower(): c for c in CHG_SERVICE_CATEGORIES}
    validated = []
    seen = set()
    for raw in raw_categories or []:
        item = str(raw).strip()
        if not item:
            continue
        key = item.lower()
        if key in canonical_by_key:
            canon = canonical_by_key[key]
        elif key.startswith("autre:") or key.startswith("autre :"):
            canon = item
        else:
            canon = f"autre: {item}"
        if canon.lower() not in seen:
            seen.add(canon.lower())
            validated.append(canon)
    return validated


def classify_expertise_categories(client, services_requis, livrables, expertises_resume, expertises_cles, sommaire_projet):
    """
    Classify the AO's required expertise against CHG's 12 service
    categories (ddm-chg.com/expertises) for a high-level go/no-go-committee
    view. Uses only already-extracted fields as input, not the raw AO text.

    Structural (unparseable JSON) failures get one retry, matching the main
    extraction call's own JSON-parse retry -- the output is small so this
    is low-probability, but the fix is cheap and the alternative (an empty
    Column E) is worse. Content failures (hallucinated/drifted category
    names) get NO retry: validate_expertise_categories() already
    deterministically corrects that case via the "autre: X" fallback --
    retrying wouldn't improve on it.

    Session 51 fix: the initial call uses temperature=0 (consistent
    classification of the same AO across repeated runs), but the JSON-parse
    retry deliberately does NOT -- unlike the main extraction call's own
    retry (which shares temperature=0 on both attempts and relies solely on
    a bigger token budget, appropriate for its much larger schema where
    truncation is the dominant real cause), this call's tiny output
    (max_tokens=1024 for a <=5-item list) makes truncation unlikely; a
    malformed/inconsistent generation is the more plausible failure here,
    so the retry keeps a non-zero-temperature "different roll" chance
    alongside the larger token budget.

    A JSON-parse failure and a non-parse API/call failure both still
    degrade to an empty list (same as a genuine zero-match AO -- Column E
    can't visually distinguish the three today, a separate, bigger
    follow-up), but each failure origin now gets its own distinct, typed
    log line so the cause is at least diagnosable after the fact.
    """
    categories_list = "\n".join(f"- {c}" for c in CHG_SERVICE_CATEGORIES)
    input_summary = f"""Résumé du projet: {sommaire_projet}

Services requis (volets professionnels/techniques): {json.dumps(services_requis, ensure_ascii=False)}
Livrables: {json.dumps(livrables, ensure_ascii=False)}
Expertises résumées: {json.dumps(expertises_resume, ensure_ascii=False)}
Expertises clés: {json.dumps(expertises_cles, ensure_ascii=False)}"""

    prompt = f"""Tu es un expert en appels d'offres pour Groupe Conseil CHG, une firme de génie-conseil québécoise organisée en 12 catégories de services:

{categories_list}

Voici les informations déjà extraites d'un appel d'offres:

{input_summary}

TÂCHE: Classe ce mandat au niveau macro selon les catégories de services CHG ci-dessus -- l'objectif est qu'un comité go/no-go voie EN UN COUP D'OEIL à quelles catégories CHG ce mandat correspond, PAS une liste détaillée d'expertises ou de tâches.

RÈGLES:
- Retourne UNIQUEMENT les catégories de la liste ci-dessus qui s'appliquent réellement à ce mandat, avec le NOM EXACT tel qu'écrit ci-dessus (respecte l'orthographe et la ponctuation exactement).
- Si le mandat nécessite une expertise qui ne correspond à AUCUNE des 12 catégories, nomme-la explicitement au format "autre: [nom court de l'expertise]" -- ne force JAMAIS une expertise hors-liste dans la catégorie la plus proche.
- Reste au niveau macro: ne retourne PAS de sous-domaines ou de tâches spécifiques (ex: "conception de ponceaux" n'est pas une catégorie -- c'est de l'"Infrastructure" ou de la "Structure, bâtiment et ouvrage d'art" selon le contexte).
- N'invente aucune catégorie et ne modifie jamais l'orthographe des 12 noms ci-dessus.
- Retourne entre 1 et 5 catégories typiquement -- un mandat très large peut en toucher plus, un mandat très ciblé peut n'en toucher qu'une seule.

Retourne UNIQUEMENT un JSON valide avec cette structure exacte:
{{"expertise_categories": ["Catégorie 1", "Catégorie 2"]}}"""

    def _call(max_tokens=1024, temperature=None):
        _kwargs = dict(
            model="claude-sonnet-4-5-20250929",
            max_tokens=max_tokens,
            messages=[{"role": "user", "content": prompt}],
        )
        if temperature is not None:
            _kwargs["temperature"] = temperature
        response = client.messages.create(**_kwargs)
        if response.stop_reason == "max_tokens":
            raise TruncatedResponseError(f"response truncated at max_tokens={max_tokens}")
        text = response.content[0].text
        json_start = text.find("{")
        json_end = text.rfind("}") + 1
        return json.loads(text[json_start:json_end])

    try:
        parsed = _call(temperature=0)
    except (json.JSONDecodeError, ValueError, IndexError, TruncatedResponseError) as _je:
        print(f"  [expertise_categories] JSON parse failed ({_je}) — retrying once")
        try:
            parsed = _call(max_tokens=2048)
        except Exception as _e2:
            print(f"  [expertise_categories] retry failed ({type(_e2).__name__}: {_e2}) — leaving list empty (could be genuine zero-match or extraction failure — check error above)")
            return []
        print(f"  [expertise_categories] retry succeeded after JSON parse failure")
    except Exception as _e:
        print(f"  [expertise_categories] API/call failure ({type(_e).__name__}: {_e}) — no retry attempted — leaving list empty (could be genuine zero-match or extraction failure)")
        return []

    return validate_expertise_categories(parsed.get("expertise_categories", []))


# ── Main analysis ──────────────────────────────────────────────────────────────

def analyze_ao_gonogo(pdf_path):
    """Analyse un AO et retourne la grille Go/No-Go complète.

    Accepte soit un chemin unique (str -- comportement historique inchangé),
    soit une liste de chemins (nouveau: AO multi-fichiers, Session 34). Le
    premier fichier de la liste est considéré comme le document principal;
    si le texte combiné dépasse MAX_AO_CHARS, la troncature affecte la FIN
    du texte combiné en priorité (donc les documents additionnels avant le
    document principal).
    """
    paths = pdf_path if isinstance(pdf_path, list) else [pdf_path]
    _failed_additional_docs = []

    if len(paths) == 1:
        print(f"Extraction du texte: {paths[0]}")
        ao_text = extract_ao_text(paths[0])
    else:
        print(f"Extraction du texte de {len(paths)} fichiers (AO multi-fichiers)")
        _parts = []
        principal_ok = False
        for _i, _p in enumerate(paths):
            _label = "DOCUMENT PRINCIPAL" if _i == 0 else f"DOCUMENT ADDITIONNEL {_i} ({os.path.basename(_p)})"
            print(f"  Extraction: {_p}")
            _t = extract_ao_text(_p)
            if _t:
                _parts.append(f"--- {_label} ---\n{_t}")
                if _i == 0:
                    principal_ok = True
            else:
                print(f"  [avertissement] aucun texte extrait de {_p}")
                if _i != 0:
                    _failed_additional_docs.append(os.path.basename(_p))
        # Session 51: an additional document failing is still non-fatal (annexes
        # are often supplementary -- the loop above continues and joins whatever
        # succeeded, unchanged). But if the PRINCIPAL document specifically failed,
        # the combined text can still be non-empty (padded by additional docs
        # alone) and silently pass the `if not ao_text` check below -- producing a
        # plausible-looking analysis grounded in the wrong document. That's worse
        # than no analysis at all, so it's raised here explicitly rather than left
        # to the generic empty-text check.
        if not principal_ok:
            raise Exception(
                f"Document principal illisible: aucun texte extrait de {paths[0]} "
                f"({len(_parts)} document(s) additionnel(s) extrait(s) avec succès)"
            )
        ao_text = "\n\n".join(_parts)
        if len(ao_text) > MAX_AO_CHARS:
            print(f"  [extraction] ATTENTION: texte combine tronque de {len(ao_text)} a {MAX_AO_CHARS} caracteres")
            ao_text = ao_text[:MAX_AO_CHARS]

    if not ao_text:
        raise Exception("Impossible d'extraire le texte du document")

    print(f"Texte extrait: {len(ao_text)} caractères")
    print("Analyse Claude en cours...")

    client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"), max_retries=5)

    today = date.today().strftime("%d %B %Y")

    # Jours fériés Québec 2025-2026
    feries_qc = """
JOURS FÉRIÉS QUÉBEC 2025-2026 (à exclure de l'échéancier):
- 2025-12-25 Noël
- 2025-12-26 Lendemain de Noël
- 2026-01-01 Jour de l'An
- 2026-01-02 Lendemain du Jour de l'An
- 2026-04-03 Vendredi saint
- 2026-05-18 Journée nationale des patriotes
- 2026-06-24 Fête nationale du Québec
- 2026-07-01 Fête du Canada
- 2026-09-07 Fête du Travail
- 2026-10-12 Action de grâces
- 2026-11-11 Jour du Souvenir (non férié QC mais noter)
- 2026-12-25 Noël
RÈGLE: Exclure TOUS les samedis, dimanches et jours fériés ci-dessus du calcul des dates.
RÈGLE PUROLATOR: Si format_depot = "physique" ou "papier" ou "en personne", ajouter +2 jours ouvrables supplémentaires pour la livraison Purolator.
"""

    prompt = f"""Tu es un expert en appels d'offres pour une firme de génie-conseil québécoise (Groupe Conseil CHG).
DATE D'AUJOURD'HUI: {today} — utilise cette date pour évaluer les délais. Si la date de dépôt est déjà passée, indique clairement que l'AO est EXPIRÉ.

{feries_qc}

DOCUMENT AO:
{ao_text}

Retourne UNIQUEMENT un JSON valide avec cette structure exacte — respecte TOUS les noms de clés exactement tels qu'écrits:
{{
  "identification": {{
    "numero_ao": "",
    "titre_projet": "",
    "client": "",
    "contact_client": "",
    "email_contact": "",
    "telephone_contact": "",
    "date_depot": "",
    "heure_depot": "",
    "format_depot": "",
    "format_depot_citation": "",
    "lieu_depot": "",
    "date_limite_questions": "",
    "date_emission": "",
    "timezone_projet": "",
    "soumission_physique": false,
    "delai_purolator_requis": false
  }},
  "description_projet": {{
    "description": "",
    "sommaire_projet": "Résumé clair en 3-4 phrases du projet: quoi, où, pourquoi, ampleur. Langage non-technique compréhensible par tous.",
    "expertises_requises_resume": ["expertise 1", "expertise 2"],
    "localisation": "",
    "region_administrative": "",
    "services_requis": [],
    "livrables": [],
    "echeancier_client": "",
    "intrants_fournis_client": [],
    "budget_estime": ""
  }},
  "specs_techniques_redaction": {{
    "nb_pages_max": "",
    "typographie": "",
    "format_requis": "",
    "nb_copies": "",
    "autres_specs": ""
  }},
  "criteres_conformite": {{
    "visite_lieux_obligatoire": "non_specifie",
    "visite_lieux_obligatoire_citation": "",
    "visite_lieux_date": null,
    "visite_lieux_heure": null,
    "visite_lieux_lieu": null,
    "garantie_soumission_requise": false,
    "montant_garantie": "",
    "assurances_requises": [],
    "montant_assurance_responsabilite": "",
    "logiciels_obligatoires": [],
    "formations_certifications_requises": [],
    "criteres_eliminatoires": []
  }},
  "criteres_evaluation": {{
    "mode_attribution": "",
    "mode_evaluation": "",
    "ponderation_technique": "",
    "ponderation_financiere": "",
    "note_minimale_qualite": "",
    "entrevue_prevue": false,
    "organigramme_requis": false,
    "nombre_repondants_attendus": "",
    "criteres_detail": [
      {{
        "numero_critere": "2.7.1",
        "intitule": "Nom du critère",
        "points_max": 20,
        "contenu_requis": "Description détaillée de ce qui est évalué",
        "nb_projets_requis": 5,
        "annees_experience_requises": 10
      }}
    ]
  }},
  "ressources_requises": {{
    "expertises_cles": [],
    "cv_requis": true,
    "fiches_projets_requises": true,
    "nombre_cv": "",
    "montant_minimal_projets": "",
    "mode_remuneration": "",
    "sous_traitance_autorisee": "non_specifie",
    "sous_traitance_autorisee_citation": "",
    "sous_traitance_recommandee": false,
    "domaines_sous_traitance": [],
    "cumul_roles_autorise": "non_specifie",
    "cumul_roles_autorise_citation": "",
    "roles_detailles": [
      {{
        "role": "",
        "titre_requis": "",
        "annees_experience_min": 0,
        "annees_experience_mentionnees": false,
        "qualifications_obligatoires": [],
        "qualifications_souhaitees": [],
        "points_differenciateurs": "",
        "profil_junior_acceptable": false,
        "nb_personnes": 1
      }}
    ]
  }},
  "incongruites": [
    {{
      "section": "",
      "description": "",
      "question_a_poser": ""
    }}
  ],
  "risques_contractuels": {{
    "type_contrat": "",
    "penalites": "",
    "clauses_responsabilite": "",
    "consortium_autorise": "non_specifie",
    "consortium_autorise_citation": "",
    "autres_disciplines_separees": "non_specifie",
    "autres_disciplines_separees_citation": "",
    "autres_risques": []
  }},
  "analyse_strategique": {{
    "historique_client_chg": "",
    "niveau_concurrence_estime": "",
    "concurrents_region": [],
    "points_forts_chg": [],
    "points_faibles_chg": [],
    "score_gonogo": 0,
    "recommandation": "",
    "justification": ""
  }},
  "echeancier_redaction": {{
    "date_depot": "",
    "date_document_final": "",
    "date_limite_ingenieur": "",
    "date_iso": "",
    "date_reunion_planification": "",
    "date_livraison_purolator": "",
    "notes_echeancier": ""
  }}
}}

RÈGLES IMPORTANTES:
- ÉTAPE 1 (à faire en PREMIER, avant toute autre analyse): compare la date_depot à la date d'aujourd'hui ({today}).
  * Si date_depot est dans le FUTUR (AO actif): NE MENTIONNE PAS la question d'expiration dans justification, ni de brouillon de raisonnement à ce sujet. Passe directement à l'ÉTAPE 2.
  * Si date_depot est PASSÉE (AO expiré): score_gonogo = 0, recommandation = "NO GO", justification = "APPEL D'OFFRES EXPIRÉ — [raison courte]". Ne fais PAS l'analyse Go/No-Go complète dans ce cas — arrête-toi ici pour ce champ.
- ÉTAPE 2 (seulement si l'AO est actif): Score Go/No-Go: 0-100 (70+ = GO, 50-69 = CONDITIONNEL, moins de 50 = NO GO).
- CRITIQUE: les champs score_gonogo et recommandation DOIVENT correspondre EXACTEMENT à ta conclusion finale. N'écris jamais une valeur provisoire (ex: liée à une fausse alerte d'expiration) que tu corriges ensuite uniquement dans le texte de justification — si tu te corriges, le champ structuré doit refléter la version corrigée, jamais la version initiale.
- Détecte TOUTES les incongruités, contradictions ou ambiguïtés dans le document

RÈGLES CRITIQUES POUR criteres_detail — RESPECTER EXACTEMENT CES NOMS DE CLÉS:
  * "numero_critere" (pas "numero", pas "critere", pas "id")
  * "intitule" (pas "titre", pas "nom", pas "name")
  * "points_max" (pas "points", pas "ponderation", pas "weight")
  * "contenu_requis" (pas "description", pas "contenu", pas "details")
  * "nb_projets_requis" (nombre entier, 0 si non spécifié)
  * "annees_experience_requises" (nombre entier, 0 si non spécifié)

RÈGLES POUR roles_detailles:
- Extraire CHAQUE rôle mentionné dans l'AO avec ses exigences exactes
- Si annees_experience non mentionnées: annees_experience_mentionnees = false, profil_junior_acceptable = true
- "points_differenciateurs": identifier si ce rôle génère des points selon un critère d'évaluation.
  Format: "Critère 2.7.X — X pts — [condition qui donne des points ex: 10 ans exp = X pts]"
  Si aucun lien direct avec un critère: laisser vide ""

RÈGLES POUR domaines_sous_traitance ET sous_traitance_recommandee:
- Si l'AO mentionne qu'une étude ou un service (ex: étude géotechnique, arpentage,
  étude environnementale, relevé topographique) a DÉJÀ été réalisé, existe déjà,
  ou sera fourni par le client/donneur d'ouvrage, NE PAS l'inclure dans
  domaines_sous_traitance — ce n'est plus un besoin de sous-traitance pour ce projet.
- N'inclure dans domaines_sous_traitance QUE les expertises manquantes que CHG
  devra réellement sous-traiter pour réaliser le mandat.

RÈGLES POUR sous_traitance_autorisee ET consortium_autorise (ET leurs champs
"_citation" associés, sous_traitance_autorisee_citation ET consortium_autorise_citation):
- Ces deux champs de valeur doivent être "oui", "non" ou "non_specifie" (jamais
  un booléen).
- OBLIGATOIRE — procédure à suivre DANS CET ORDRE, séparément pour chacun des
  deux sujets (sous-traitance, consortium) :
  1. Cherche dans le texte de l'AO une phrase qui mentionne EXPLICITEMENT ce
     sujet précis (le mot "sous-traitance"/"sous-traitant" pour le premier
     champ, le mot "consortium" pour le second).
  2. Si tu trouves une telle phrase: recopie-la (ou l'extrait pertinent, une
     phrase max) telle quelle dans le champ "_citation" correspondant, puis
     détermine "oui" ou "non" à partir de CETTE phrase précise (autorise vs.
     interdit/encadre restrictivement).
  3. Si tu NE trouves AUCUNE phrase mentionnant explicitement ce sujet: le
     champ "_citation" reste "" (chaîne vide) ET le champ de valeur DOIT être
     "non_specifie". Il est INTERDIT d'écrire "non" quand "_citation" est
     vide — l'absence de citation possible EST la preuve qu'il faut utiliser
     "non_specifie", jamais une raison de déduire "non".
- Chaque sujet est INDÉPENDANT: le fait que la sous-traitance soit autorisée
  explicitement (ou interdite) ne dit RIEN sur le consortium, et vice-versa —
  ne déduis JAMAIS l'un de l'autre par contraste ou par analogie. Un AO peut
  très bien encadrer la sous-traitance sans jamais aborder les consortiums:
  dans ce cas, consortium_autorise = "non_specifie", même si
  sous_traitance_autorisee = "oui" ou "non".
- "Non" et "non spécifié" ne sont PAS équivalents : "Non" signifie que l'AO
  interdit ou encadre restrictivement la pratique EXPLICITEMENT, pas que le
  sujet est absent du document.

RÈGLES POUR cumul_roles_autorise (ET son champ "_citation" associé,
cumul_roles_autorise_citation) -- même structure que sous_traitance_autorisee/
consortium_autorise ci-dessus:
- Ce champ de valeur doit être "oui", "non" ou "non_specifie" (jamais un
  booléen).
- OBLIGATOIRE -- procédure à suivre DANS CET ORDRE :
  1. Cherche dans le texte de l'AO une phrase qui mentionne EXPLICITEMENT
     qu'une même personne peut (ou ne peut pas) occuper/cumuler plus d'un
     rôle principal ou plus d'une fonction dans le cadre du contrat.
  2. Si tu trouves une telle phrase : recopie-la (ou l'extrait pertinent,
     une phrase max) telle quelle dans le champ "_citation", puis détermine
     "oui" (le cumul est explicitement permis, ex: "une même personne peut
     cumuler ces fonctions") ou "non" (le cumul est explicitement interdit,
     ex: "doit être occupé par une ressource distincte", "ne peut être
     affectée à plus d'un rôle principal") à partir de CETTE phrase précise.
  3. Si tu NE trouves AUCUNE phrase mentionnant explicitement le cumul de
     rôles/fonctions (au sens de l'étape 1) : le champ "_citation" reste ""
     (chaîne vide) ET le champ de valeur DOIT être "non_specifie". Il est
     INTERDIT d'écrire "non" quand "_citation" est vide.
- EXCLUS EXPLICITEMENT de cette détection (ne traite PAS ces mentions comme
  des citations de cumul de rôles, même si elles emploient un vocabulaire
  proche) -- si la SEULE mention trouvée dans le document est de l'un de ces
  deux types, traite-la comme une absence de mention explicite (retombe à
  l'étape 3 ci-dessus) :
  * les clauses de relève/remplacement en cas d'absence (ex: "une même
    personne peut être la relève de plus d'un membre de l'équipe") -- c'est
    une règle de continuité de service en cas d'absence, pas une permission
    de cumul de rôles ;
  * les exigences de continuité du personnel-clé (ex: interdiction de
    remplacer le chargé de projet pendant la période d'évaluation des
    soumissions, ou sans consentement écrit après adjudication) -- c'est une
    règle de stabilité de l'équipe dans le temps, pas une déclaration sur le
    cumul de rôles.
- "Non" et "non_specifie" ne sont PAS équivalents : "Non" signifie que l'AO
  interdit explicitement le cumul de rôles ; "non_specifie" signifie que le
  sujet est totalement absent du document, ou que seule une mention exclue
  ci-dessus (relève/continuité) a été trouvée.

RÈGLES POUR autres_disciplines_separees (ET son champ "_citation" associé,
autres_disciplines_separees_citation):
- Ce champ de valeur doit être "oui", "non" ou "non_specifie" (jamais un
  booléen).
- OBLIGATOIRE -- procédure à suivre DANS CET ORDRE :
  1. Cherche dans le texte de l'AO une phrase qui mentionne EXPLICITEMENT
     qu'une discipline ou profession PRÉCISE ET NOMMÉE (ex: architecture,
     géotechnique, arpentage, électricité) sera mandatée ou procurée
     directement par le donneur d'ordre au moyen d'un contrat séparé/
     distinct de celui visé par le présent appel d'offres.
  2. Si tu trouves une telle phrase : recopie-la (ou l'extrait pertinent,
     une phrase max) telle quelle dans le champ "_citation", puis détermine
     "oui" à partir de CETTE phrase précise.
  3. Si tu NE trouves AUCUNE phrase nommant une discipline/profession
     précise procurée séparément par le donneur d'ordre : le champ
     "_citation" reste "" (chaîne vide) ET le champ de valeur DOIT être
     "non_specifie".
- EXIGENCE CRITIQUE -- discipline NOMMÉE obligatoire pour "oui" : une clause
  générique de coordination avec "tout tiers" ou "toute Personne" ayant un
  contrat distinct avec le donneur d'ordre, SANS nommer la discipline ou
  profession visée (ex: une clause type "Contrats simultanés" imposant à
  l'adjudicataire de collaborer avec quiconque a un contrat séparé avec le
  donneur d'ordre, sans préciser qui), N'EST PAS une discipline procurée
  séparément -- c'est une clause standard de coordination multi-contrats.
  Dans ce cas : le champ "_citation" reste "" (chaîne vide) ET le champ de
  valeur DOIT être "non_specifie" (le sujet précis -- une discipline nommée
  procurée séparément -- n'est pas explicitement abordé), jamais "oui".
- "Non" et "non_specifie" ne sont PAS équivalents : "Non" signifie que l'AO
  aborde explicitement le sujet et indique qu'aucune autre discipline n'est
  procurée séparément ; "non_specifie" signifie que le sujet est totalement
  absent du document, ou que seule une clause générique de coordination
  multi-contrats (sans discipline nommée) a été trouvée.

RÈGLES POUR visite_lieux_obligatoire (ET son champ "_citation" associé,
visite_lieux_obligatoire_citation) -- même structure que sous_traitance_autorisee/
consortium_autorise ci-dessus:
- Ce champ de valeur doit être "oui", "non" ou "non_specifie" (jamais un
  booléen).
- OBLIGATOIRE -- procédure à suivre DANS CET ORDRE :
  1. Cherche dans le texte de l'AO une phrase qui mentionne EXPLICITEMENT
     une visite des lieux (visite de chantier, visite du site, visite
     obligatoire ou optionnelle, etc.).
  2. Si tu trouves une telle phrase : recopie-la (ou l'extrait pertinent,
     une phrase max) telle quelle dans le champ "_citation", puis détermine
     "oui" (visite obligatoire) ou "non" (visite explicitement optionnelle,
     ou explicitement aucune visite prévue) à partir de CETTE phrase précise.
  3. Si tu NE trouves AUCUNE phrase mentionnant une visite des lieux : le
     champ "_citation" reste "" (chaîne vide) ET le champ de valeur DOIT
     être "non_specifie". Il est INTERDIT d'écrire "non" quand "_citation"
     est vide.
- "Non" et "non_specifie" ne sont PAS équivalents : "Non" signifie que l'AO
  aborde explicitement le sujet et indique que la visite n'est pas
  obligatoire ; "non_specifie" signifie que le sujet est totalement absent
  du document.

RÈGLES POUR visite_lieux_date, visite_lieux_heure, visite_lieux_lieu:
- Champs indépendants de visite_lieux_obligatoire -- extraire ces logistiques
  QUELLE QUE SOIT la valeur de visite_lieux_obligatoire ("oui", "non" ou
  "non_specifie"), si l'AO les mentionne explicitement (ex: une visite
  optionnelle peut quand même avoir une date/heure/lieu précis).
- N'extraire QUE ce qui est explicitement indiqué dans le texte. Ne JAMAIS
  déduire, estimer ou inventer une valeur manquante.
- visite_lieux_date : format ISO "AAAA-MM-JJ" (ex: "2026-09-09"). null si
  aucune date n'est mentionnée.
- visite_lieux_heure : l'heure telle qu'indiquée (ex: "11h00"). null si
  aucune heure n'est mentionnée.
- visite_lieux_lieu : l'adresse ou le lieu de rencontre tel qu'indiqué. null
  si aucun lieu n'est mentionné.

RÈGLES POUR intrants_fournis_client:
- Extraire la liste des éléments, études, données ou documents que le client
  (Ville/donneur d'ouvrage) fournit DÉJÀ ou mettra à la disposition de
  l'adjudicataire (ex: relevés existants, études géotechniques antérieures,
  plans, données SIG, autorisations déjà obtenues).
- Chercher spécifiquement les sections du type « Intrants fournis par la
  Ville », « Documents fournis », « Informations disponibles ».
- Format: liste de courtes entrées (une par intrant), sans détail exhaustif
  — ex: "Étude géotechnique complémentaire (2024)", "Plans d'arpentage existants".
- Si aucun intrant n'est mentionné: liste vide [].

RÈGLES POUR specs_techniques_redaction:
- Extraire le nombre de pages maximum si mentionné dans l'AO
- Extraire la typographie requise (police, taille) si mentionnée
- Extraire le nombre de copies requises
- Si non spécifié: laisser la valeur vide ""

RÈGLES POUR sommaire_projet:
- 3 à 4 phrases maximum
- Langage clair, non-technique
- Inclure: quoi (type de travaux), où (localisation), ampleur (budget/distance/superficie), contexte si pertinent

RÈGLES POUR services_requis:
- Lister les VOLETS PROFESSIONNELS ET TECHNIQUES du mandat, de façon EXHAUSTIVE:
  les grandes catégories de travail technique/professionnel que l'adjudicataire
  doit réaliser (ex: conception architecturale, mécanique, électrique; ingénierie
  civile; arpentage; relevés; études géotechniques; surveillance des travaux;
  supervision de la mise en route), incluant les volets techniques spécifiques
  qui n'apparaissent que dans le détail du mandat (ex: installation d'un
  équipement particulier, raccordement à un système existant, démantèlement
  d'installations, gestion des sols contaminés, etc.).
- EXCLURE de cette liste le contenu administratif/procédural du cycle
  contractuel — ce contenu appartient EXCLUSIVEMENT au champ "livrables",
  jamais à "services_requis". Exemples à exclure: demandes de permis et
  autorisations, préparation des documents d'appel d'offres (rédaction,
  avis, addendas), analyse des soumissions et recommandation, décomptes
  progressifs et recommandations de paiement, gestion des déficiences,
  comptes rendus de réunions, attestations de conformité, plans comme
  construits.
- Ne pas énumérer ces tâches administratives même si l'AO les détaille dans
  une section "Biens livrables" ou équivalente — cette section alimente le
  champ "livrables", pas "services_requis".
- Chaque volet professionnel/technique distinct doit apparaître UNE SEULE
  FOIS, peu importe le niveau de détail auquel l'AO le mentionne (résumé en
  introduction OU détail plus loin dans le document) — ne pas relister la
  même activité deux fois à des niveaux de granularité différents.

RÈGLES POUR livrables:
- Lister les ARTEFACTS ET DOCUMENTS CONCRETS que l'adjudicataire doit
  produire et remettre (plans, devis, rapports, attestations, comptes
  rendus, décomptes, etc.).
- C'est ICI, et non dans "services_requis", que doivent apparaître les
  tâches administratives et procédurales du cycle contractuel: demandes de
  permis et autorisations, documents d'appel d'offres (avis, addendas),
  analyse des soumissions, décomptes progressifs, gestion des déficiences,
  comptes rendus de réunions, attestations de conformité, plans comme
  construits.

RÈGLES POUR expertises_requises_resume ET expertises_cles:

CE QUI COMPTE COMME EXPERTISE (à inclure):
- UNIQUEMENT des spécialisations professionnelles ou corps de métier requis pour
  répondre à l'AO. Teste chaque candidat avec la question: "est-ce que ce terme
  répond à 'quelle est ma spécialisation professionnelle?' (= expertise, à
  inclure) ou à 'que fais-tu concrètement dans ce mandat?' (= activité/tâche, à
  exclure)?"
- Catégories courantes, à titre d'exemples SEULEMENT — liste illustrative et
  NON EXHAUSTIVE, ne pas s'y limiter et ne pas forcer un terme du document dans
  une de ces catégories s'il décrit une spécialisation plus précise ou
  différente: ingénierie civile municipale, hydraulique, génie géotechnique,
  génie mécanique, génie électrique, génie structural, architecture,
  architecture du paysage, arpentage, environnement, biologie/écologie,
  géomatique.

CE QUI N'EST PAS UNE EXPERTISE (à exclure, même si le document AO l'appelle
"expertise" ou la liste dans une section de ce type):
- Les tâches, activités, livrables ou fonctions de gestion — ce que
  l'adjudicataire FAIT, pas ce qu'il EST. Exemples INCORRECTS (ce sont des
  tâches, pas des spécialisations): "Coordination intervenants", "Gestion de
  projet" (fonction de gestion, pas une spécialisation technique),
  "Surveillance de chantier" / "Surveillance des travaux" (une activité de
  surveillance, pas une spécialisation en soi).
- Les connaissances de normes, règlements ou certifications ne sont pas des
  expertises non plus — si le domaine technique sous-jacent est déjà couvert
  par un autre terme de la liste, ne pas ajouter la norme/certification comme
  entrée séparée.

CONSOLIDATION — chaque expertise distincte doit apparaître UNE SEULE FOIS:
- Avant de finaliser la liste, compare CHAQUE paire d'éléments entre eux: si
  l'un est une reformulation ou une variante du NOM d'un autre — le même
  professionnel, avec la même compétence, décrit avec des mots différents —
  fusionne-les en gardant la formulation la plus précise et complète, et
  supprime le doublon.
- Test pour décider si deux termes désignent VRAIMENT la même spécialisation:
  un professionnel qualifié pour l'un serait-il automatiquement qualifié
  pour l'autre, et les deux termes seraient-ils interchangeables dans une
  description de poste ? Si OUI → fusionner. Si NON → garder les deux entrées
  séparées, même si elles appartiennent au même domaine général.
- IMPORTANT — ne PAS fusionner un terme spécifique dans une catégorie plus
  large sous prétexte qu'il en fait partie: "être une sous-catégorie ou une
  composante de X" n'est PAS la même chose que "être une reformulation de
  X". Une spécialisation technique précise qui requiert des compétences
  distinctes de sa catégorie générale doit rester une entrée séparée, même
  si elle est mentionnée dans le même domaine général qu'une autre entrée.
- Exemples réels de fusions CORRECTES (même spécialisation, mots
  différents) — à appliquer par analogie, liste NON exhaustive:
  * "Ingénierie civile municipale", "Ingénierie civile", "Gestion de projets
    d'infrastructures" et "Gestion de projets d'infrastructures municipales"
    désignent la MÊME spécialisation → garder UNE SEULE entrée: "Ingénierie
    civile municipale".
  * "Ingénierie de la mobilité et transport" et "Ingénierie de la mobilité"
    désignent la MÊME spécialisation → garder UNE SEULE entrée (la plus
    complète): "Ingénierie de la mobilité et transport".
- Exemples réels de NON-FUSIONS (spécialisations distinctes dans un même
  domaine général — à ne PAS fusionner entre elles NI absorber dans une
  catégorie plus large):
  * "Ingénierie routière et conception géométrique", "Structures de
    chaussée", "Signalisation routière" et "Expertise en pistes cyclables
    et aménagements actifs" sont QUATRE spécialisations distinctes du génie
    routier — ne PAS les fusionner entre elles, et ne PAS les absorber dans
    "Ingénierie civile municipale" au motif qu'elles en font partie. Chacune
    exige des compétences techniques différentes (géométrie et tracé /
    dimensionnement de chaussée / normes de signalisation / aménagements
    cyclables) qu'un généraliste en ingénierie civile municipale ne possède
    pas nécessairement — garder les CINQ entrées séparées.

RÈGLES POUR mode_evaluation:
- Classifier le mode d'évaluation en une des 3 valeurs EXACTES: "Prix seul" | "Qualité seule" | "Qualité et prix"
- "Prix seul": attribution au plus bas soumissionnaire conforme, sans évaluation de la qualité
- "Qualité seule": évaluation uniquement sur critères qualitatifs, sans facteur prix dans la formule d'attribution
- "Qualité et prix": la formule d'attribution combine un pointage qualité ET le prix soumis (ex: système à pointage intérimaire + prix, ratio prix/pointage, etc.)
- Se baser sur le texte de la section "Évaluation des offres" / "Adjudication du contrat" de l'AO, notamment le champ mode_attribution.

RÈGLES POUR concurrents_region:
- Laisser [] — sera rempli via analyse historique CHG (Phase 2)

RÈGLES POUR format_depot, format_depot_citation, soumission_physique ET
delai_purolator_requis:
- Une mention de "SEAO" (ou de tout autre portail similaire) N'IMPORTE OÙ
  dans le document N'EST PAS une preuve, dans un sens ou dans l'autre, que
  le dépôt électronique de la soumission est autorisé. SEAO sert dans la
  grande majorité des AO UNIQUEMENT à obtenir les documents d'appel
  d'offres et à consulter les addendas — ce rôle NE dit RIEN sur le mode de
  dépôt de la soumission elle-même.
- OBLIGATOIRE — procédure à suivre DANS CET ORDRE:
  1. Localise la section qui décrit spécifiquement la RÉCEPTION ou la
     PRÉSENTATION de la soumission elle-même (titres typiques: "Réception
     des soumissions", "Dépôt et ouverture des soumissions", "Présentation
     de la soumission", "Présentation sur le SEAO"). Ignore les sections
     qui parlent seulement d'obtention des documents d'appel d'offres,
     d'addendas, ou de vérification préalable auprès du SEAO — ces
     sections NE déterminent PAS le mode de dépôt de la soumission.
  2. Recopie, dans "format_depot_citation", la phrase (un extrait pertinent,
     une phrase max) de CETTE section précise qui décrit comment la
     soumission doit être remise.
  3. Détermine format_depot À PARTIR de cette citation UNIQUEMENT:
     - Si cette section décrit UNIQUEMENT une remise physique (dépôt à une
       adresse, dans une boîte/chute, enveloppe scellée remise en personne
       ou par courrier) SANS y mentionner d'option de transmission
       électronique de la soumission elle-même: format_depot = "Physique"
       — même si SEAO est mentionné ailleurs dans le document pour
       l'obtention des documents ou des addendas.
     - Si cette section précise décrit EXPLICITEMENT une option de
       transmission électronique de la soumission elle-même (ex: "les
       soumissions peuvent être transmises par voie électronique via le
       SEAO", "présenter sa soumission sur le SEAO", "déposer via
       [plateforme]"): format_depot doit refléter cette option électronique
       (ex: "Physique ou électronique (SEAO)").
  4. Si aucune phrase de cette section précise ne mentionne explicitement
     une option électronique de dépôt de la soumission: format_depot_citation
     reste "" (vide) ET format_depot ne doit PAS contenir "électronique".
- CHG dépose TOUJOURS par voie électronique lorsque cette section précise
  l'autorise, même si elle permet aussi (ou privilégie) le dépôt physique.
  Ne jamais choisir soumission_physique = true simplement parce que le
  dépôt physique est mentionné ou est l'option par défaut de l'AO.
- soumission_physique = false ET delai_purolator_requis = false UNIQUEMENT
  si la section de réception/présentation des soumissions autorise
  explicitement un dépôt électronique de la soumission elle-même — jamais
  sur la seule base d'une mention de SEAO ou d'un autre portail ailleurs
  dans le document.
- soumission_physique = true SI la section de réception/présentation des
  soumissions n'autorise QUE le dépôt physique (aucune option électronique
  de dépôt de la soumission elle-même décrite dans cette section précise).

RÈGLES POUR l'échéancier (calcul à rebours, weekends ET jours fériés exclus):
  * date_reunion_planification: dès que possible après émission
  * date_document_final: 4 jours ouvrables avant dépôt (2e étape après réunion)
  * date_limite_ingenieur: 6 jours ouvrables avant dépôt
  * date_iso: 2 jours ouvrables avant dépôt
  * date_livraison_purolator: si soumission physique, date_depot - 2 jours ouvrables
  * Si soumission physique: intégrer +2 jours ouvrables pour Purolator dans tous les calculs

- Pour timezone_projet: noter si le projet est dans un fuseau différent de EST
- Réponds UNIQUEMENT avec le JSON, sans texte avant ou après, sans backticks markdown"""

    def _call_claude_and_parse(max_tokens=8000, substitute=False):
        _prompt = apply_refusal_substitutions(prompt) if substitute else prompt
        response = client.messages.create(
            model="claude-sonnet-4-5-20250929",
            max_tokens=max_tokens,
            temperature=0,
            messages=[{"role": "user", "content": _prompt}]
        )
        if response.stop_reason == "max_tokens":
            raise TruncatedResponseError(f"response truncated at max_tokens={max_tokens}")
        if response.stop_reason == "refusal" or not response.content:
            raise ResponseRefusedError(f"main extraction response refused or empty (stop_reason={response.stop_reason!r})")
        result_text = response.content[0].text.strip()
        if result_text.startswith("```"):
            result_text = result_text.split("\n", 1)[1].rsplit("```", 1)[0].strip()
        return json.loads(result_text)

    # Session 52: the ResponseRefusedError catch wraps BOTH the initial call
    # and the truncation-retry below, not just the first -- confirmed via a
    # real reprocessing of AO 13203's full 4-document text (300k chars,
    # hitting MAX_AO_CHARS) that the truncation-retry itself can *also* come
    # back refused, not just truncated again. A single-level try (catching
    # ResponseRefusedError only around the first call) would let that second
    # refusal propagate uncaught. The substitution retry uses max_tokens=
    # 16000 to match whichever budget was actually in play when the refusal
    # happened, so it isn't wasted on a response that gets truncated too.
    try:
        try:
            result = _call_claude_and_parse()
        except (json.JSONDecodeError, IndexError, TruncatedResponseError) as _je:
            print(f"  [JSON parse retry] initial parse failed ({_je}) — retrying once with higher max_tokens")
            result = _call_claude_and_parse(max_tokens=16000)
    except ResponseRefusedError as _re:
        print(f"  [refusal retry] call refused ({_re}) — retrying once with known-trigger-phrase substitution")
        result = _call_claude_and_parse(max_tokens=16000, substitute=True)

    # ── Completeness check: criteres_detail points should sum close to 100 ──
    # Extraction variance observed in Session 23 (5 vs 3 criteres_detail
    # entries on identical input). temperature=0 reduces this but a document
    # can still be genuinely ambiguous — retry once if the sum looks short.
    #
    # Session 51: arithmetic runs on a local normalize_criteres_detail()
    # copy (never persisted -- Post-processor 1 below remains the sole
    # authoritative write into result) so non-dict/non-numeric entries
    # (Session 50's Column 9 case) can't crash this block and silently
    # skip the retry. The retry call itself is caught in its own narrow
    # try/except so a transient API/parse failure on the retry can't wipe
    # out an already-successful original extraction.
    try:
        def _attempt_retry():
            # Session 51: shared one-shot retry-call wrapper for both
            # triggers below. Returns the fresh result on success, or None
            # if the retry call itself failed -- caller keeps result as-is.
            try:
                return _call_claude_and_parse()
            except (anthropic.APIError, IndexError, AttributeError, json.JSONDecodeError) as _retry_e:
                print(f"  [completeness check] retry call failed ({_retry_e}) — keeping original extraction, retry skipped")
                return None

        _ce = result.get("criteres_evaluation")
        if not isinstance(_ce, dict):
            # Session 51: criteres_evaluation itself malformed (e.g. null) --
            # same retry-once treatment as a low points sum below. No reset
            # is written here either way; Post-processor 1 is the sole
            # persist point if this retry doesn't produce a dict either.
            print(f"  [completeness check] criteres_evaluation malformed (type={type(_ce).__name__}) — retrying extraction once")
            _retried = _attempt_retry()
            if _retried is not None:
                result = _retried
        else:
            _criteres_check = normalize_criteres_detail(_ce.get("criteres_detail", []))
            _points_sum = sum(c.get("points_max", 0) for c in _criteres_check)
            if _criteres_check and _points_sum < 90:
                print(f"  [completeness check] criteres_detail points sum to {_points_sum}/100 — retrying extraction once")
                _retried = _attempt_retry()
                if _retried is not None:
                    result = _retried
                    _criteres_recheck = normalize_criteres_detail(result.get("criteres_evaluation", {}).get("criteres_detail", []))
                    _points_recheck = sum(c.get("points_max", 0) for c in _criteres_recheck)
                    print(f"  [completeness check] retry result: {_points_recheck}/100 across {len(_criteres_recheck)} criteres")
    except (KeyError, TypeError) as _e:
        print(f"  [completeness check] SKIPPED (unexpected result shape): {_e}")

    # ── Post-processor 1: normalize criteres_detail keys (schema enforcement) ──
    try:
        if not isinstance(result.get("criteres_evaluation"), dict):
            # Session 51: still malformed after the completeness check's own
            # retry attempt (or that retry wasn't triggered/available) --
            # this is the sole place that persists a reset, so every
            # downstream consumer always gets a well-shaped dict.
            print(f"  [criteres normalization] criteres_evaluation still malformed — resetting to empty (sole persist point)")
            result["criteres_evaluation"] = {"criteres_detail": []}
        else:
            criteres_raw = result["criteres_evaluation"].get("criteres_detail", [])
            if criteres_raw:
                result["criteres_evaluation"]["criteres_detail"] = normalize_criteres_detail(criteres_raw)
                print(f"  criteres_detail normalisés: {len(result['criteres_evaluation']['criteres_detail'])} critères")
    except Exception as e:
        print(f"  [criteres normalization] avertissement: {e}")

    # ── Post-processor 2: expiry detection (Python is authoritative, not Claude) ──
    try:
        from datetime import date as _date
        import re as _re

        _MOIS_FR = {
            "janvier":1,"février":2,"fevrier":2,"mars":3,"avril":4,
            "mai":5,"juin":6,"juillet":7,"août":8,"aout":8,
            "septembre":9,"octobre":10,"novembre":11,"décembre":12,"decembre":12
        }
        _MOIS_EN = {
            "january":1,"february":2,"march":3,"april":4,"may":5,"june":6,
            "july":7,"august":8,"september":9,"october":10,"november":11,"december":12
        }

        def _parse_date(s):
            """Parse FR / EN / ISO date strings. Returns date or None."""
            if not s:
                return None
            s = s.strip()
            # ISO: 2025-11-25
            m = _re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
            if m:
                return _date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            # French: 25 novembre 2025
            m = _re.match(r'(\d{1,2})\s+([a-zA-Zéûôùàâè]+)\s+(\d{4})', s, _re.IGNORECASE)
            if m:
                mois = m.group(2).lower()
                num  = _MOIS_FR.get(mois) or _MOIS_EN.get(mois)
                if num:
                    return _date(int(m.group(3)), num, int(m.group(1)))
            # English: November 25, 2025
            m = _re.match(r'([a-zA-Z]+)\s+(\d{1,2})[,\s]+(\d{4})', s, _re.IGNORECASE)
            if m:
                num = _MOIS_EN.get(m.group(1).lower())
                if num:
                    return _date(int(m.group(3)), num, int(m.group(2)))
            return None

        _date_depot_str = result.get("identification", {}).get("date_depot", "")
        _parsed = _parse_date(_date_depot_str)
        if _date_depot_str and not _parsed:
            print(f"  [expiry check] date_depot present but unparseable ({_date_depot_str!r}) — expiry override skipped, Claude's own recommendation stands unmodified")

        if _parsed and _parsed < _date.today():
            _strat = result.get("analyse_strategique", {})
            # Preserve analytical score and recommendation for reference
            _strat["score_analytique"]         = _strat.get("score_gonogo", 0)
            _strat["recommandation_analytique"] = _strat.get("recommandation", "—")
            # Override to EXPIRÉ
            _strat["score_gonogo"]    = 0
            _strat["recommandation"]  = "EXPIRÉ"
            _strat["justification"]   = (
                f"APPEL D'OFFRES EXPIRÉ — La date limite de dépôt était le "
                f"{_date_depot_str}. Cet appel d'offres ne peut plus être soumissionné. "
                f"Score analytique conservé à titre de référence : "
                f"{_strat['score_analytique']}/100 — {_strat['recommandation_analytique']}."
            )
            result["analyse_strategique"] = _strat
    except Exception as _e:
        print(f"  [expiry check] avertissement: {_e}")

    # ── Post-processor 3: bureau_recommande (haversine distance to CHG offices) ──
    try:
        OFFICES_CHG = {
            "Quebec":   (46.806, -71.247),
            "Montreal": (45.550, -73.629),
            "Saguenay": (48.427, -71.067),
        }

        # Static proxy-city coordinates for Quebec's 17 official administrative
        # regions, used only as a fallback when localisation itself can't be
        # geocoded. Approximate city-center coordinates -- adequate precision
        # for choosing among 3 offices spread across the province.
        REGION_PROXY_COORDS = {
            "Bas-Saint-Laurent": (48.4489, -68.5222),
            "Saguenay–Lac-Saint-Jean": (48.4279, -71.0685),
            "Capitale-Nationale": (46.8139, -71.2082),
            "Mauricie": (46.3432, -72.5432),
            "Estrie": (45.4042, -71.8929),
            "Montreal": (45.5019, -73.5674),
            "Montréal": (45.5019, -73.5674),
            "Outaouais": (45.4765, -75.7013),
            "Abitibi-Témiscamingue": (48.2436, -79.0234),
            "Côte-Nord": (49.2158, -68.1517),
            "Nord-du-Québec": (49.9169, -74.3699),
            "Gaspésie–Îles-de-la-Madeleine": (48.8347, -64.4861),
            "Chaudière-Appalaches": (46.7382, -71.2492),
            "Laval": (45.6066, -73.7124),
            "Lanaudière": (46.0157, -73.4441),
            "Laurentides": (45.7811, -74.0037),
            "Montérégie": (45.5312, -73.5183),
            "Centre-du-Québec": (45.8837, -72.4842),
        }

        def _haversine_km(coord_a, coord_b):
            lat1, lon1 = coord_a
            lat2, lon2 = coord_b
            r = 6371.0
            p1, p2 = math.radians(lat1), math.radians(lat2)
            dphi = math.radians(lat2 - lat1)
            dlambda = math.radians(lon2 - lon1)
            a = (math.sin(dphi / 2) ** 2
                 + math.cos(p1) * math.cos(p2) * math.sin(dlambda / 2) ** 2)
            return 2 * r * math.asin(math.sqrt(a))

        def _geocode(query):
            """Returns (lat, lon) or None. Never raises."""
            if not query or not query.strip():
                return None
            try:
                resp = requests.get(
                    "https://nominatim.openstreetmap.org/search",
                    params={"q": f"{query}, Québec, Canada", "format": "json", "limit": 1},
                    headers={"User-Agent": "CHG-AO-Pipeline/1.0 (contact: patrick.salvano@gmail.com)"},
                    timeout=8,
                )
                data = resp.json()
                if data:
                    return (float(data[0]["lat"]), float(data[0]["lon"]))
            except Exception as _geo_e:
                print(f"  [bureau_recommande] geocode failed for '{query}': {_geo_e}")
            return None

        _localisation = result.get("description_projet", {}).get("localisation", "")
        _region = result.get("description_projet", {}).get("region_administrative", "")

        def _geocode_progressive(localisation_str):
            """
            Tries the full localisation string first; if Nominatim returns
            nothing (common with street-level intersections like
            "Rues X et Y, Ville, Region"), progressively strips the leading
            comma-separated segment (most-specific/street-level detail) and
            retries, until a match is found or nothing usable remains.
            Returns (point, methode) -- methode is "geocode_localisation" for
            the full string, or "geocode_localisation_partiel" for a
            stripped-down retry.
            """
            if not localisation_str or not localisation_str.strip():
                return None, None
            parts = [p.strip() for p in localisation_str.split(",") if p.strip()]
            if not parts:
                return None, None
            pt = _geocode(localisation_str)
            if pt:
                return pt, "geocode_localisation"
            for i in range(1, len(parts)):
                candidate = ", ".join(parts[i:])
                pt = _geocode(candidate)
                if pt:
                    return pt, "geocode_localisation_partiel"
            return None, None

        _point, _methode = _geocode_progressive(_localisation)
        if not _methode:
            _methode = "geocode_localisation"

        if not _point and _region:
            _point = REGION_PROXY_COORDS.get(_region.strip())
            _methode = "centroide_region" if _point else "aucun"
        elif not _point:
            _methode = "aucun"

        if _point:
            _distances = {
                office: round(_haversine_km(_point, coords), 1)
                for office, coords in OFFICES_CHG.items()
            }
            _bureau_recommande = min(_distances, key=_distances.get)
            result["description_projet"]["bureau_recommande"] = _bureau_recommande
            result["description_projet"]["bureau_recommande_methode"] = _methode
            result["description_projet"]["distances_bureau_km"] = _distances
            print(f"  [bureau_recommande] {_bureau_recommande} (méthode: {_methode}, distances: {_distances})")
        else:
            result["description_projet"]["bureau_recommande"] = None
            result["description_projet"]["bureau_recommande_methode"] = "aucun"
            result["description_projet"]["distances_bureau_km"] = {}
            print("  [bureau_recommande] aucune localisation exploitable -- champ laissé à null")
    except Exception as _e:
        print(f"  [bureau_recommande] avertissement: {_e}")
        # Session 52: guarantee the three enrichment keys always exist with a
        # value distinguishable from the normal "aucun localisation exploitable"
        # completion state -- a mid-block crash previously left them entirely
        # unset, indistinguishable downstream from "ran fine, found nothing."
        if not isinstance(result.get("description_projet"), dict):
            result["description_projet"] = {}
        result["description_projet"].setdefault("bureau_recommande", None)
        result["description_projet"].setdefault("bureau_recommande_methode", "erreur")
        result["description_projet"].setdefault("distances_bureau_km", {})

    # ── Post-processor 4: keyword safety net for consortium_autorise / sous_traitance_autorisee ──
    # Session 48: code-level backstop for the prompt rule above. If Claude
    # returns "oui"/"non" but the AO source text never literally mentions the
    # topic, the value can't be evidence-based -- downgrade to non_specifie.
    try:
        # autres_disciplines_separees needs more than "contrat distinct/séparé"
        # present ANYWHERE in the document -- that phrase alone also appears in
        # generic multi-prime-contractor coordination boilerplate that names no
        # discipline (real example: AO-26-20-P-SP's "10.05 Contrats simultanés"
        # clause). A named discipline/profession must co-occur near the
        # contract-split language for this to be verifiable evidence of a
        # specific discipline procured separately -- list scoped to disciplines
        # realistically named as separate from CHG's own 12 service categories
        # (Infrastructure, Structure, Hydrologie/hydraulique, Arpentage/drone/
        # bathymétrie, Transport, Gestion de projet, Environnement, Études
        # économiques, Foresterie, Géomatique/TI, Planification stratégique,
        # Urbanisme) -- most commonly architecture on building-envelope AOs,
        # plus other building/specialty disciplines CHG doesn't itself offer.
        _NAMED_DISCIPLINE_RE = re.compile(
            r"architecture|architecte|paysag(?:iste|er)|"
            r"g[ée]otechni(?:que|cien)|"
            r"hydrog[ée]olog(?:ie|ique)|"
            r"arpentage|arpenteur|"
            r"[ée]lectrom[ée]canique|"
            r"[ée]lectricit[ée]|g[ée]nie\s+[ée]lectrique|"
            r"m[ée]canique(?:\s+du\s+b[âa]timent)?|g[ée]nie\s+m[ée]canique|"
            r"biologie|[ée]cologie|"
            r"acoustique|"
            r"(?:protection|s[ée]curit[ée])\s+incendie|"
            r"laboratoire|contr[ôo]le\s+des\s+mat[ée]riaux",
            re.IGNORECASE,
        )
        _CONTRACT_SPLIT_RE = re.compile(
            r"contrat\s+(?:distinct|s[eé]par[eé])|mandat[eé]s?\s+directement|"
            r"(?:autre|distinct|s[eé]par[eé])\s+mode\s+d['’]attribution|"
            r"mode\s+d['’]attribution\s+(?:distinct|s[eé]par[eé])",
            re.IGNORECASE,
        )

        def _autres_disciplines_evidence(text, window=300):
            for _m in _CONTRACT_SPLIT_RE.finditer(text):
                _start = max(0, _m.start() - window)
                _end = min(len(text), _m.end() + window)
                if _NAMED_DISCIPLINE_RE.search(text[_start:_end]):
                    return True
            return False

        _SAFETY_NET_FIELDS = [
            ("risques_contractuels", "consortium_autorise", re.compile(r"consortium|co[-\s‑]*entreprise|regroupement\s+de\s+personnes", re.IGNORECASE)),
            ("ressources_requises", "sous_traitance_autorisee", re.compile(r"sous[-\s‑]*traitan", re.IGNORECASE)),
            ("ressources_requises", "cumul_roles_autorise", re.compile(r"cumul|plus\s+d['’]un\s+r[oô]le|plus\s+d['’]une\s+fonction", re.IGNORECASE)),
            ("risques_contractuels", "autres_disciplines_separees", _autres_disciplines_evidence),
        ]
        for _section, _field, _checker in _SAFETY_NET_FIELDS:
            _val = result.get(_section, {}).get(_field)
            _found = _checker(ao_text) if callable(_checker) else _checker.search(ao_text)
            if _val in ("oui", "non") and not _found:
                print(f"  [safety_net] {_field}: downgraded from '{_val}' to 'non_specifie' -- keyword not found in source AO text")
                result[_section][_field] = "non_specifie"
                # Session 54: autres_disciplines_separees ONLY -- also clear the
                # stale citation on downgrade, so "non_specifie" never ships next
                # to leftover citation text (contradicts the prompt's own
                # "non_specifie implies empty citation" invariant). consortium_
                # autorise/sous_traitance_autorisee have the identical gap but are
                # deliberately left untouched here -- logged as a separate backlog
                # item for a future dedicated session, not fixed opportunistically
                # alongside this one.
                if _field == "autres_disciplines_separees":
                    result[_section][_field + "_citation"] = ""
    except Exception as _e:
        print(f"  [safety_net] avertissement: {_e}")

    # ── Post-processor 5: classify expertise against CHG's 12 service categories (Column E reframe) ──
    try:
        _descr = result.get("description_projet", {})
        _ress = result.get("ressources_requises", {})
        result.setdefault("description_projet", {})["expertise_categories"] = classify_expertise_categories(
            client,
            _descr.get("services_requis", []),
            _descr.get("livrables", []),
            _descr.get("expertises_requises_resume", []),
            _ress.get("expertises_cles", []),
            _descr.get("sommaire_projet", ""),
        )
        print(f"  expertise_categories classifiées: {result['description_projet']['expertise_categories']}")
    except Exception as _e:
        print(f"  [expertise_categories] avertissement: {_e}")
        result.setdefault("description_projet", {})["expertise_categories"] = []
    # ── Fin post-processors ──────────────────────────────────────────────────

    # Session 52: surface any non-principal document extraction failure into
    # the returned result -- previously a print-only signal fully discarded
    # once this function returned (see pipeline_w1.py's degraded_reasons wiring).
    result["documents_extraction_failed"] = _failed_additional_docs

    # Session 52: single, unconditional restoration point -- covers the main
    # extraction call's own fields AND Post-processor 5's
    # classify_expertise_categories() output (which reads from `result`
    # above, before this point) in one pass. Safe/idempotent whether or not
    # the refusal-retry path was ever taken this run.
    result = restore_refusal_substitutions(result)

    print("Analyse Go/No-Go complétée!")
    return result


if __name__ == "__main__":
    import sys
    pdf_path = sys.argv[1] if len(sys.argv) > 1 else "/root/test_ao.pdf"

    result = analyze_ao_gonogo(pdf_path)

    output_path = pdf_path.rsplit(".", 1)[0] + "_gonogo.json"
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\nRésultat sauvegardé: {output_path}")
    print(f"\nScore Go/No-Go: {result['analyse_strategique']['score_gonogo']}/100")
    print(f"Recommandation: {result['analyse_strategique']['recommandation']}")
    print(f"Incongruités détectées: {len(result['incongruites'])}")
    print(f"Rôles détaillés: {len(result['ressources_requises']['roles_detailles'])}")
    criteres = result.get('criteres_evaluation', {}).get('criteres_detail', [])
    print(f"Critères d'évaluation: {len(criteres)}")
    specs = result.get('specs_techniques_redaction', {})
    if any(specs.values()):
        print(f"Specs rédaction: pages max={specs.get('nb_pages_max','—')}, typo={specs.get('typographie','—')}")
