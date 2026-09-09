import os, sys, json, io, tempfile, datetime, re
import anthropic
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from googleapiclient.discovery import build
from google.oauth2 import service_account
from googleapiclient.http import MediaIoBaseDownload
sys.path.insert(0, '/root')
from analyze_ao import (
    TruncatedResponseError, ResponseRefusedError,
    apply_refusal_substitutions, restore_refusal_substitutions,
)

# ── Args from n8n ─────────────────────────────────────────────────────────────
FILE_ID   = sys.argv[1] if len(sys.argv) > 1 else None
FILE_NAME = sys.argv[2] if len(sys.argv) > 2 else "AO_inconnu"
print(f"🚀 Pipeline W2 (Agent 3) started for: {FILE_NAME} ({FILE_ID})")

SERVICE_ACCOUNT_FILE = '/root/chg-credentials/service-account.json'
creds = service_account.Credentials.from_service_account_file(
    SERVICE_ACCOUNT_FILE, scopes=['https://www.googleapis.com/auth/drive'])
drive = build('drive', 'v3', credentials=creds)
claude = anthropic.Anthropic()

# ── Load libraries ────────────────────────────────────────────────────────────
cvs_raw  = json.load(open('/root/chg-library/cv_docx_library.json'))
cvs      = [{'nom': cv.get('nom',''), 'titre': cv.get('nom',''), 'competences': [], 'text': cv.get('texte','')} for cv in cvs_raw]
projects = json.load(open('/root/chg-library/projects_library.json'))
real_proposal = json.load(open('/root/chg-library/real_proposal_analysis.json'))
crossref = json.load(open('/root/chg-library/ao_os_crossref.json'))
crossref_intel = {'lecons': [], 'arguments_differenciateurs': [], 'mots_cles_techniques': [], 'mappings': []}
for pair in crossref.get('pairs', []):
    crossref_intel['lecons'].extend(pair.get('lecons_pour_futurs_aos', []))
    crossref_intel['arguments_differenciateurs'].extend(pair.get('arguments_differenciateurs', []))
    crossref_intel['mots_cles_techniques'].extend(pair.get('mots_cles_techniques', []))
    crossref_intel['mappings'].extend(pair.get('comment_chg_repond_aux_exigences', []))
crossref_intel['lecons'] = list(dict.fromkeys(crossref_intel['lecons']))[:20]
crossref_intel['arguments_differenciateurs'] = list(dict.fromkeys(crossref_intel['arguments_differenciateurs']))[:10]
crossref_intel['mots_cles_techniques'] = list(dict.fromkeys(crossref_intel['mots_cles_techniques']))[:15]

# ── Helpers ───────────────────────────────────────────────────────────────────
def download_bytes(file_id):
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, drive.files().get_media(fileId=file_id))
    done = False
    while not done: _, done = dl.next_chunk()
    return buf.getvalue()

def claude_call(prompt, max_tokens=2000):
    r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}])
    resp = r.content[0].text.strip()
    if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
    return resp

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        run.font.bold = True

def add_divider(doc):
    p = doc.add_paragraph('—' * 80)
    p.runs[0].font.size = Pt(6)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

def match_cv(member_name, cvs):
    member_lower = member_name.lower()
    member_base  = member_lower.split(' - ')[0].strip()
    member_words = set(member_base.split())
    best_cv, best_score = None, 0
    for c in cvs:
        cv_base  = c['nom'].lower().split(' - ')[0].strip()
        cv_words = set(cv_base.split())
        score    = len(member_words & cv_words)
        if score > best_score:
            best_score = score
            best_cv    = c
    return best_cv if best_score >= 1 else None

def get_sdt_value(cell):
    """Extract the selected value from a Word content control (dropdown) in a cell."""
    try:
        from lxml import etree
        tc = cell._tc
        sdts = tc.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdt')
        for sdt in sdts:
            content = sdt.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdtContent')
            if content is not None:
                texts = []
                for elem in content.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                    if elem.text:
                        texts.append(elem.text)
                val = ''.join(texts).strip()
                if val:
                    return val
    except Exception:
        pass
    return cell.text.strip()

# ── STEP 1: Download and parse approved planning doc ──────────────────────────
print("📄 Downloading approved planning document...")
planning_bytes = download_bytes(FILE_ID)
planning_doc   = Document(io.BytesIO(planning_bytes))

# Derive AO name from filename: strip _approved suffix and extensions
safe_name = FILE_NAME
for ext in ['.pdf', '.docx', '.doc']:
    safe_name = safe_name.replace(ext, '').replace(ext.upper(), '')
safe_name = safe_name.replace('_approved', '').replace('PLANIFICATION_', '')
safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', safe_name)[:60]

# ── STEP 2: Extract confirmed resources from planning doc ─────────────────────
print("🔍 Extracting confirmed resources and fiches from planning doc...")
confirmed_team   = []  # [{nom, role, selection}]
confirmed_fiches = []  # [nom_fiche]
fiche_routing = {}  # {nom_fiche: 'firme'|'charge_projet'|'les_deux'|'exclure'}
banque_selections = []  # [{nom, role_dans_projet, selection}]
extra_instructions = ""

tables = planning_doc.tables
for tbl_idx, tbl in enumerate(tables):
    rows = tbl.rows
    if not rows:
        continue
    # Detect resources table by header content
    header_text = ' '.join([c.text.strip() for c in rows[0].cells]).lower()
    if 'ressource' in header_text and ('rôle' in header_text or 'role' in header_text):
        print(f"  → Found resources table (table {tbl_idx})")
        for row in rows[1:]:
            cells = row.cells
            if len(cells) < 3:
                continue
            role       = cells[0].text.strip()
            nom        = cells[1].text.strip()
            if not nom or not role:
                continue
            # Last cell is the dropdown
            selection = get_sdt_value(cells[-1])
            if not selection:
                selection = cells[-1].text.strip()
            print(f"    {nom} | {role} | {selection}")
            if True:  # Store all — excluded members needed for replacement logic
                confirmed_team.append({
                    'nom': nom,
                    'role_dans_projet': role,
                    'selection': selection  # Principale or Relève
                })
    # Detect fiches table
    elif 'fiche' in header_text or 'projet' in header_text:
        print(f"  \u2192 Found fiches table (table {tbl_idx})")
        for row in rows[1:]:
            cells = row.cells
            if len(cells) < 2:
                continue
            nom_fiche = cells[0].text.strip()
            if not nom_fiche:
                continue
            selection = get_sdt_value(cells[-1])
            if not selection:
                selection = cells[-1].text.strip()
            sel_norm = selection.strip().lower()
            print(f"    {nom_fiche} | {selection}")
            if sel_norm in ('firme', 'firm'):
                routage = 'firme'
            elif sel_norm in ('charg\u00e9 de projet', 'charge de projet', 'cdp'):
                routage = 'charge_projet'
            elif sel_norm in ('les deux', 'les 2', 'both'):
                routage = 'les_deux'
            elif sel_norm in ('exclure', 'non', 'no'):
                routage = 'exclure'
            else:
                # Unknown/blank value -> safe default, never silently drop a fiche
                routage = 'les_deux'
            fiche_routing[nom_fiche] = routage
            if routage != 'exclure':
                confirmed_fiches.append(nom_fiche)
    elif 'nom' in header_text and 'selectionner' in header_text:
        # Detect Banque de ressources disponibles table
        print(f'  -> Found banque table (table {tbl_idx})')
        for row in rows[1:]:
            cells = row.cells
            if len(cells) < 3:
                continue
            nom_banque = cells[0].text.strip()
            role_banque = cells[1].text.strip()
            if not nom_banque:
                continue
            selection = get_sdt_value(cells[-1])
            if not selection:
                selection = cells[-1].text.strip()
            print(f'     {nom_banque} | {selection}')
            if selection.lower() in ['oui', 'yes']:
                banque_selections.append({'nom': nom_banque, 'role_dans_projet': role_banque, 'selection': 'Principale'})

# Extract extra instructions from the free-text box
# It's a 1x1 table with cream background — look for it by process of elimination
for tbl in tables:
    rows = tbl.rows
    if len(rows) >= 1 and len(tbl.columns) == 1:
        text = '\n'.join([r.cells[0].text.strip() for r in rows if r.cells[0].text.strip()])
        # Exclude if it's the resources or fiches table
        if text and 'ressource' not in text.lower() and 'fiche' not in text.lower():
            placeholder_phrases = ['rédacteur', 'alexandra', 'inscrire', 'client connu', 'préciser']
            if not any(ph in text.lower() for ph in placeholder_phrases):
                extra_instructions = text
                print(f"  → Extra instructions found: {text[:100]}...")

# Add banque selections to confirmed team
for b in banque_selections:
    if not any(m['nom'] == b['nom'] for m in confirmed_team):
        confirmed_team.append(b)
        print(f'  + Banque ajout: {b["nom"]} ({b["role_dans_projet"]})')
print(f"✅ Confirmed team: {[m['nom'] for m in confirmed_team]}")
print(f"✅ Confirmed fiches: {confirmed_fiches}")

# Fallback: if parsing extracted nothing, use pipeline.py's normal selection
if not confirmed_team:
    print("⚠️ No team extracted from planning doc — falling back to AI selection")

# ── STEP 3: Find the original AO (same name without _approved/PLANIFICATION) ──
# Search Drive for the original AO PDF in CHG - Appels d'Offres folder
AO_FOLDER_ID = '1cXGe0n7GRSzRP1KlmcR6siDI8lppEH48'
print(f"🔍 Searching for original AO in Drive (safe_name: {safe_name})...")
ao_text = ""
ao_data = {}
ao_data_parse_failed = False
ao_pdf_path = None

try:
    results = drive.files().list(
        q=f"'{AO_FOLDER_ID}' in parents and trashed=false",
        fields="files(id,name)").execute()
    ao_files = results.get('files', [])
    # Find best match by comparing safe_name tokens
    safe_tokens = set(safe_name.lower().replace('_',' ').split())
    best_match, best_score = None, 0
    for f in ao_files:
        fn = re.sub(r'[^a-zA-Z0-9 ]', ' ', f['name'].lower())
        fn_tokens = set(fn.split())
        score = len(safe_tokens & fn_tokens)
        if score > best_score:
            best_score = score
            best_match = f
    if best_match and best_score >= 2:
        print(f"✅ Found AO: {best_match['name']} (score {best_score})")
        ao_bytes = download_bytes(best_match['id'])
        is_docx  = best_match['name'].lower().endswith(('.docx', '.doc'))
        if is_docx:
            ao_pdf_tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
            ao_pdf_tmp.write(ao_bytes); ao_pdf_tmp.close()
            ao_pdf_path = ao_pdf_tmp.name
            ao_doc_tmp  = Document(io.BytesIO(ao_bytes))
            ao_text     = '\n'.join([p.text for p in ao_doc_tmp.paragraphs if p.text.strip()])[:10000]
        else:
            ao_pdf_tmp = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
            ao_pdf_tmp.write(ao_bytes); ao_pdf_tmp.close()
            ao_pdf_path = ao_pdf_tmp.name
            from pdfminer.high_level import extract_text as _et
            ao_text = _et(ao_pdf_path)[:10000]
    else:
        print(f"⚠️ Could not find original AO in Drive (best score: {best_score}) — proceeding without AO text")
except Exception as e:
    print(f"⚠️ AO search failed: {e}")

# ── STEP 4: Analyze AO text to get ao_data ────────────────────────────────────
if ao_text:
    today = datetime.date.today().strftime('%Y-%m-%d')

    def _step4_call(text):
        _r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=8000,
            messages=[{"role": "user", "content": f"Aujourd'hui: {today}\nAO:\n{text}\n\nUNIQUEMENT JSON sans backticks:\n{{\"client\":\"\",\"numero_ao\":\"\",\"titre_projet\":\"\",\"date_limite\":\"\",\"type_travaux\":\"\",\"competences_requises\":[],\"localisation\":\"\",\"resume_mandat\":\"\",\"points_addendas\":[]}}"}])
        # Session 52: same undefended empty-content/refusal crash pattern as
        # pipeline_w1.py's STEP 2 had before today (r.content[0] on an empty
        # list) -- confirmed via the same real production incident (AO
        # 13203), since this call re-extracts and re-sends the same
        # principal-document text.
        if not _r.content:
            if _r.stop_reason == "refusal":
                raise ResponseRefusedError(
                    f"STEP 4: Claude refused to respond (stop_reason=refusal) — "
                    f"content policy trigger, not a parse/truncation issue."
                )
            raise ResponseRefusedError(f"STEP 4: Claude response has no content (stop_reason={_r.stop_reason!r})")
        return _r

    try:
        r = _step4_call(ao_text)
    except ResponseRefusedError as _re:
        print(f"  [refusal retry] STEP 4 initial call refused ({_re}) — retrying once with known-trigger-phrase substitution")
        r = _step4_call(apply_refusal_substitutions(ao_text))

    resp = r.content[0].text.strip()
    if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
    try:
        ao_data = json.loads(resp)
    except Exception as _json_e:
        print(f"⚠️ STEP 4 JSON parse failed ({_json_e}) — falling back to filename-derived ao_data")
        ao_data = {"titre_projet": safe_name, "client": "", "type_travaux": ""}
        ao_data_parse_failed = True
else:
    ao_data = {"titre_projet": safe_name, "client": "", "type_travaux": ""}

print(f"✅ AO: {ao_data.get('client','')} — {ao_data.get('titre_projet','')}")

# ── STEP 5: Build team_data from confirmed selections ─────────────────────────
if confirmed_team:
    team_data = {"equipe": [{"nom": m["nom"], "role_dans_projet": m["role_dans_projet"]} for m in confirmed_team if m.get("selection","").lower() not in ["exclure", "exclude"]]}
else:
    # Fallback: AI selection
    cv_summary = [{"nom": c["nom"], "titre": c.get("titre",""), "competences": c.get("competences",[])} for c in cvs]
    r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=800,
        messages=[{"role": "user", "content": f"AO: {json.dumps(ao_data, ensure_ascii=False)}\nRessources: {json.dumps(cv_summary, ensure_ascii=False)}\nSélectionne 5-6 ressources. UNIQUEMENT JSON sans backticks:\n{{\"equipe\":[{{\"nom\":\"\",\"role_dans_projet\":\"\"}}]}}"}])
    resp = r.content[0].text.strip()
    if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
    team_data = json.loads(resp)

print(f"✅ Team: {[m['nom'] for m in team_data['equipe']]}")

# ── STEP 5b: Handle excluded resources + find replacements ───────────────────
excluded_members = [m for m in confirmed_team if m.get('selection','').lower() in ['exclure', 'exclude']]
if excluded_members:
    print(f"🔄 {len(excluded_members)} resource(s) excluded — searching for replacements...")
    excluded_roles = [m['nom'] + ' (' + m['role_dans_projet'] + ')' for m in excluded_members]
    cv_summary = [{"nom": c["nom"], "titre": c.get("titre",""), "texte_resume": c.get("texte","")[:300]} for c in cvs]
    confirmed_names_lower = [m['nom'].lower() for m in team_data['equipe']]
    cv_candidates = [c for c in cv_summary if c['nom'].lower() not in confirmed_names_lower]
    try:
        repl_prompt = "Alexandra a exclu ces ressources: " + str(excluded_roles) + "\n"
        repl_prompt += "Instructions d'Alexandra: " + extra_instructions + "\n"
        repl_prompt += "AO type: " + ao_data.get('type_travaux','') + "\n"
        repl_prompt += "CVs disponibles: " + json.dumps(cv_candidates[:20], ensure_ascii=False) + "\n"
        repl_prompt += "Pour chaque rôle exclu, propose UN remplaçant. UNIQUEMENT JSON sans backticks:\n"
        repl_prompt += '{"remplacements": [{"role": "", "nom_remplacant": "", "justification": ""}]}'
        r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=1000,
            messages=[{"role": "user", "content": repl_prompt}])
        resp = r.content[0].text.strip()
        if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
        repl_data = json.loads(resp)
        for repl in repl_data.get('remplacements', []):
            nom = repl.get('nom_remplacant','')
            role = repl.get('role','')
            if nom and nom.lower() not in confirmed_names_lower:
                team_data['equipe'].append({"nom": nom, "role_dans_projet": role})
                print("✅ Replacement added: " + nom + " → " + role)
    except Exception as e:
        print("⚠️ Replacement search failed: " + str(e))

# ── STEP 5c: Infer bureau for each team member ────────────────────────────────
try:
    team_cv_info = []
    for m in team_data['equipe']:
        cv = match_cv(m['nom'], cvs)
        cv_text = cv.get('texte', cv.get('text', ''))[:800] if cv else ''
        team_cv_info.append({"nom": m['nom'], "cv_extrait": cv_text})
    bureau_prompt = "Pour chaque membre, détermine son bureau (Québec, Montréal, ou Saguenay) selon son CV.\n"
    bureau_prompt += "Membres: " + json.dumps(team_cv_info, ensure_ascii=False) + "\n"
    bureau_prompt += 'UNIQUEMENT JSON sans backticks:\n{"bureaux": [{"nom": "", "bureau": ""}]}'
    r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=500,
        messages=[{"role": "user", "content": bureau_prompt}])
    resp = r.content[0].text.strip()
    if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
    bureau_data = json.loads(resp)
    bureau_map = {b['nom']: b['bureau'] for b in bureau_data.get('bureaux', [])}
    for m in team_data['equipe']:
        m['bureau'] = bureau_map.get(m['nom'], '')
    print("✅ Bureaux: " + str(bureau_map))
except Exception as e:
    print("⚠️ Bureau inference failed: " + str(e))

# ── STEP 6: Build confirmed project list ─────────────────────────────────────
def parse_project_regex(proj):
    import re
    texte    = proj.get('texte', '')
    nom      = proj.get('nom', proj.get('fichier', 'Projet'))
    nom_fiche = proj.get('nom', '')
    lines    = [l.strip() for l in texte.split('\n') if l.strip() and len(l.strip()) > 5]
    nom_projet = lines[0][:80] if lines else nom_fiche
    client = ''
    for line in lines:
        if 'client' in line.lower() and '|' in line:
            parts = line.split('|')
            for i, p in enumerate(parts):
                if 'client' in p.lower() and i+1 < len(parts):
                    client = parts[i+1].strip()[:60]
                    break
    m = re.search(r'(?:Municipalit[eé]|Ville|MRC|Office|Soci[eé]t[eé]|Commission|[A-Z][a-z]+ de [A-Z])["\n]{0,3,40}', texte)
    if not client and m: client = m.group(0).strip()[:60]
    m = re.search(r'(\d{2}-\d{3,}|\d{4}-\d{3,})', nom_fiche)
    annee = f"20{m.group(1)[:2]}" if m else ''
    m = re.search(r'\d{2,3}[kK$M]|\d[\d\s]{2,8}\$', texte)
    valeur = m.group(0).strip() if m else ''
    m = re.search(r'Description du projet["\n]{0,2}([\s\S]{10,200})', texte, re.IGNORECASE)
    services = m.group(1).strip()[:200] if m else texte[:200]
    return {'nom_projet': nom_projet, 'client': client, 'annee': annee, 'valeur': valeur, 'services_rendus': services}

if confirmed_fiches:
    selected_projects = []
    for nom_fiche in confirmed_fiches:
        for p in projects:
            if nom_fiche.lower() in p.get('nom','').lower() or p.get('nom','').lower() in nom_fiche.lower():
                p2 = dict(p)
                p2['_routage'] = fiche_routing.get(nom_fiche, 'les_deux')
                selected_projects.append(p2)
                break
    if not selected_projects:
        selected_projects = projects[:5]
else:
    # Fallback: keyword-based selection from ao_data
    ao_text_lower = (ao_data.get('type_travaux','') + ' ' + ao_data.get('titre_projet','')).lower()
    keywords = re.findall(r'\b\w{4,}\b', ao_text_lower)
    scored = []
    for p in projects:
        texte = p.get('texte','').lower()
        nom   = p.get('nom','').lower()
        score = sum(1 for kw in keywords if kw in texte or kw in nom)
        scored.append((score, p))
    scored.sort(key=lambda x: -x[0])
    selected_projects = [p for _, p in scored[:5]]

parsed_projects = [parse_project_regex(p) for p in selected_projects]
for _idx, _p in enumerate(selected_projects):
    if _idx < len(parsed_projects):
        parsed_projects[_idx]['routage'] = _p.get('_routage', 'les_deux')

# Enrich with Planification fiche data (similarites, pertinence, critere_couvert)
try:
    _planif_json_path = planning_json_path if 'planning_json_path' in dir() else None
    if not _planif_json_path:
        import glob as _glob
        _planif_files = _glob.glob('/local-files/PLANIFICATION_*.json')
        _safe = safe_name if 'safe_name' in dir() else ''
        _match = [f for f in _planif_files if _safe and _safe[:10].lower() in f.lower()]
        _planif_json_path = _match[0] if _match else (_planif_files[-1] if _planif_files else None)
    if _planif_json_path:
        import os as _os
        if _os.path.exists(_planif_json_path):
            _planif = json.load(open(_planif_json_path, encoding='utf-8'))
            _planif_fiches = _planif.get('fiches_suggerees', _planif.get('fiches', []))
            _fiche_map = {f.get('nom_fiche','').lower(): f for f in _planif_fiches}
            for proj in parsed_projects:
                _key = proj.get('nom_projet','').lower()
                _match_fiche = None
                for _fk, _fv in _fiche_map.items():
                    if _key[:20] in _fk or _fk[:20] in _key:
                        _match_fiche = _fv
                        break
                if _match_fiche:
                    proj['similarites']     = _match_fiche.get('similarites', '')
                    proj['pertinence_rich'] = _match_fiche.get('pertinence', '')
                    proj['critere_couvert'] = _match_fiche.get('critere_couvert', '')
except Exception as _e:
    print(f"  [planif enrich] warning: {_e}")

print(f"✅ Projects: {[p['nom_projet'][:40] for p in parsed_projects]}")


# ── STEP 7: Detect AO evaluation grid ────────────────────────────────────────
print("🔍 Detecting AO evaluation grid...")
eval_grid = []
if ao_text:
    try:
        r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=1000,
            messages=[{"role": "user", "content": f"AO:\n{ao_text}\n\nDétecte les sections d'évaluation avec leurs pondérations. UNIQUEMENT JSON sans backticks:\n{{\"sections\":[{{\"numero\":\"\",\"titre\":\"\",\"ponderation\":0,\"description\":\"\"}}]}}"}])
        resp = r.content[0].text.strip()
        if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
        eval_grid = json.loads(resp).get('sections', [])
        print(f"✅ Eval grid: {len(eval_grid)} sections")
    except Exception as e:
        print(f"⚠️ Eval grid detection failed: {e}")

if not eval_grid:
    eval_grid = [
        {"numero": "1", "titre": "Expérience du soumissionnaire", "ponderation": 25, "description": "Profil de la firme et mandats similaires"},
        {"numero": "2", "titre": "Expérience du chargé de projet", "ponderation": 30, "description": "Présentation, disponibilité, projets pertinents"},
        {"numero": "3", "titre": "Expérience des autres membres", "ponderation": 20, "description": "Équipe et capacité de relève"},
        {"numero": "4", "titre": "Compréhension du mandat et méthodologie", "ponderation": 25, "description": "Approche, échéancier, moyens"},
    ]

# ── STEP 8: Generate all OS sections ─────────────────────────────────────────
print("✍️ Generating OS sections...")
style_guide = real_proposal.get('style_guide', real_proposal.get('synthese', str(real_proposal)[:2000]))
extra_str   = f"\n\nINSTRUCTIONS SUPPLÉMENTAIRES DE LA PLANIFICATION:\n{extra_instructions}" if extra_instructions else ""

# Build team summary like pipeline.py
charge_projet = team_data['equipe'][0] if team_data['equipe'] else {"nom": "Charles Gauthier", "role_dans_projet": "Chargé de projet"}
charge_cv = match_cv(charge_projet['nom'], cvs)
charge_cv_text = charge_cv.get('texte', charge_cv.get('text', ''))[:2000] if charge_cv else ""
autres_membres = team_data['equipe'][1:] if len(team_data['equipe']) > 1 else []
autres_cvs_summary = []
for m in autres_membres:
    cv = match_cv(m['nom'], cvs)
    if cv:
        autres_cvs_summary.append(f"{m['nom']} ({m['role_dans_projet']}): {cv.get('texte', cv.get('text',''))[:500]}")

projects_summary = json.dumps([{
    'nom': p.get('nom_projet', p.get('nom','')),
    'client': p.get('client',''),
    'annee': p.get('annee',''),
    'valeur': p.get('valeur_contrat',''),
    'services': p.get('services_rendus', p.get('texte',''))[:200]
} for p in parsed_projects], ensure_ascii=False)

sections_list = "\n".join([f"  Section {s['numero']} ({s.get('poids_pct', s.get('ponderation',''))}%): {s['titre']}" for s in eval_grid])

# Build flat JSON keys like pipeline.py — proven reliable
eval_sections = eval_grid
flat_keys = {}
for s in eval_sections:
    key = f"section_{s['numero']}_{s['titre'].lower().replace(' ','_').replace('é','e').replace('è','e').replace('ê','e').replace('à','a').replace('ç','c')[:30]}"
    flat_keys[key] = f"texte complet section {s['numero']} — {s['titre']} (min 300 mots)"
flat_keys_str = json.dumps(flat_keys, ensure_ascii=False)

content_prompt = f"""Tu es un expert senior en génie-conseil québécois qui rédige des offres de service gagnantes pour le Groupe Conseil CHG inc.
TON OBJECTIF : Produire un OS qui GAGNE — pas juste qui répond à l'AO. Chaque section doit maximiser le score selon la grille d'évaluation.

=== APPEL D'OFFRES ===
{json.dumps(ao_data, ensure_ascii=False)}

=== GRILLE D'ÉVALUATION (sections scorées par l'évaluateur) ===
{sections_list}

=== CHARGÉ DE PROJET ===
{charge_projet['nom']} — {charge_projet['role_dans_projet']}
CV: {charge_cv_text}

=== AUTRES MEMBRES DE L'ÉQUIPE ===
{chr(10).join(autres_cvs_summary[:4])}

=== PROJETS DE RÉFÉRENCE CONFIRMÉS PAR ALEXANDRA ===
{projects_summary}

=== STYLE D'ÉCRITURE CHG ===
{style_guide[:1500]}

=== INTELLIGENCE STRATÉGIQUE ===
Leçons: {json.dumps(crossref_intel['lecons'][:8], ensure_ascii=False)}
Arguments différenciateurs: {json.dumps(crossref_intel['arguments_differenciateurs'][:6], ensure_ascii=False)}{extra_str}

=== INSTRUCTIONS ===
1. Écris comme un ingénieur CHG expérimenté — professionnel, précis, confiant sans arrogance
2. Chaque section doit répondre DIRECTEMENT aux critères de la grille d'évaluation
3. Section expérience firme: contrats similaires, taille firme, certifications ISO 9001:2015
4. Section chargé de projet: formation, OIQ, projets pertinents, disponibilité explicite
5. Section équipe: chaque membre avec rôle spécifique + expérience + capacité de relève
6. Section méthodologie: phases détaillées adaptées à CE projet spécifique
7. Intègre les arguments différenciateurs naturellement
8. NE PAS utiliser de langage générique ou formules creuses
9. Chaque section: minimum 300 mots, maximum 500 mots

RETOURNE UNIQUEMENT un JSON valide sans backticks:
{{
  "lettre_presentation": "lettre formelle 4 paragraphes, ton CHG, adressée au client spécifique",
  {flat_keys_str[1:-1]},
  "conclusion": "mot de la fin engageant 100 mots style CHG"
}}"""

try:
    r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=8000,
        messages=[{"role": "user", "content": content_prompt}])
    # Session 52: diagnostic-only -- already gracefully degrades to fully
    # empty content_data via the broad except below (sections_ok/
    # degraded_reasons already flag this correctly, Item 3). No retry/
    # substitution: expensive/slow call to retry blindly, no evidence this
    # specific call has hit either failure mode in production.
    if r.stop_reason == "max_tokens":
        raise TruncatedResponseError(f"STEP 8 response truncated at max_tokens=8000")
    if r.stop_reason == "refusal" or not r.content:
        raise ResponseRefusedError(f"STEP 8 response refused or empty (stop_reason={r.stop_reason!r})")
    resp = r.content[0].text.strip()
    if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
    raw_data = json.loads(resp)
    print(f"✅ Sections generated: {len(raw_data)} keys")
    # Convert flat keys to sections list format for document builder
    content_data = {
        "lettre_presentation": raw_data.get("lettre_presentation", ""),
        "conclusion": raw_data.get("conclusion", ""),
        "sections": []
    }
    for s in eval_sections:
        key = f"section_{s['numero']}_{s['titre'].lower().replace(' ','_').replace('é','e').replace('è','e').replace('ê','e').replace('à','a').replace('ç','c')[:30]}"
        contenu = raw_data.get(key, "")
        if contenu:
            content_data["sections"].append({
                "numero": s["numero"],
                "titre": s["titre"],
                "contenu": contenu
            })
    print(f"✅ Sections mapped: {len(content_data['sections'])}")
except Exception as e:
    print(f"⚠️ Section generation failed: {e}")
    content_data = {"lettre_presentation": "", "sections": [], "conclusion": ""}

# Session 51: signal for N8N_OUTPUT -- doc.save() below runs unconditionally
# regardless of content quality, so unlike pipeline_w1.py's *_report_path
# there's no existing None-style proxy for "did generation actually work".
sections_ok = bool(content_data.get("sections"))

# Session 52: single, unconditional restoration point -- covers both STEP 4's
# ao_data and STEP 8's content_data (built from ao_data) in one pass, right
# before either is used to write actual text into the document below. Safe/
# idempotent whether or not the refusal-retry path was ever taken this run.
ao_data = restore_refusal_substitutions(ao_data)
content_data = restore_refusal_substitutions(content_data)

# ── STEP 9: Build Word document ───────────────────────────────────────────────
print("📝 Building Word document...")
doc = Document()

# Page margins
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
section = doc.sections[0]
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(3)
section.right_margin  = Cm(2.5)

CHG_BLUE = RGBColor(0x1F, 0x49, 0x7D)

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = CHG_BLUE
        run.font.bold = True

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("OFFRE DE SERVICE")
run.font.size = Pt(24); run.font.bold = True; run.font.color.rgb = CHG_BLUE

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(ao_data.get('titre_projet', safe_name))
run.font.size = Pt(16); run.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(ao_data.get('client', ''))
run.font.size = Pt(14)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(f"Groupe Conseil CHG inc.\n{datetime.date.today().strftime('%d %B %Y')}")
run.font.size = Pt(12)

doc.add_page_break()

# Confidentiality clause
add_heading(doc, "Clause de confidentialité", level=2)
client_name = ao_data.get('client', 'le client')
doc.add_paragraph(
    f"Cette offre de service est la propriété du Groupe Conseil CHG. Il y est fait état du savoir-faire "
    f"de l'entreprise et des informations confidentielles qui ne peuvent être divulguées à des tiers sans "
    f"le consentement écrit préalable du Groupe Conseil CHG. Ce document a été préparé exclusivement à "
    f"l'intention de {client_name} et ne peut être utilisé à d'autres fins."
)
doc.add_page_break()

# Lettre de présentation
add_heading(doc, "Lettre de présentation")
if content_data.get('lettre_presentation'):
    for para in content_data['lettre_presentation'].split('\n\n'):
        if para.strip():
            doc.add_paragraph(para.strip())
doc.add_page_break()

def _render_project_table(doc, proj):
    add_heading(doc, proj['nom_projet'][:80], level=2)
    tbl = doc.add_table(rows=4, cols=2)
    tbl.style = 'Table Grid'
    labels = ['Ann\u00e9e', 'Pertinence', "Similarit\u00e9s avec l'AO", 'Crit\u00e8re couvert']
    vals   = [proj.get('annee',''), proj.get('pertinence_rich', proj.get('pertinence','\u00c9lev\u00e9e')),
              proj.get('similarites', proj.get('services_rendus',''))[:400], proj.get('critere_couvert','')]
    for i, (lbl, val) in enumerate(zip(labels, vals)):
        row = tbl.rows[i]
        row.cells[0].text = lbl
        row.cells[1].text = val
        row.cells[0].paragraphs[0].runs[0].font.bold = True
    doc.add_paragraph()

_routed_project_ids = set()
FIRME_KEYS = ['soumissionnaire', 'exp\u00e9rience de la firme', 'firme']
CDP_KEYS = ['charg\u00e9 de projet', 'chef de projet', 'charge de projet']

# Main sections following eval grid
for sect in content_data.get('sections', []):
    titre = f"Section {sect.get('numero','')}: {sect.get('titre','')}"
    add_heading(doc, titre, level=1)
    contenu = sect.get('contenu', '')
    for para in contenu.split('\n\n'):
        if para.strip():
            doc.add_paragraph(para.strip())

    titre_lower = sect.get('titre', '').lower()
    is_firme_section = any(k in titre_lower for k in FIRME_KEYS)
    is_cdp_section = any(k in titre_lower for k in CDP_KEYS)

    if is_firme_section or is_cdp_section:
        matched = []
        for pidx, proj in enumerate(parsed_projects):
            routage = proj.get('routage', 'les_deux')
            if routage == 'exclure':
                continue
            if is_firme_section and routage in ('firme', 'les_deux'):
                matched.append((pidx, proj))
            elif is_cdp_section and routage in ('charge_projet', 'les_deux'):
                matched.append((pidx, proj))
        if matched:
            add_heading(doc, "Projets de r\u00e9f\u00e9rence pertinents", level=2)
            for pidx, proj in matched:
                _render_project_table(doc, proj)
                _routed_project_ids.add(pidx)

    doc.add_page_break()

# Reference projects section -- fallback for any fiche not routed to a
# detected Firme/Charge de projet section title. Never silently drops
# a confirmed fiche.
_leftover = [p for i, p in enumerate(parsed_projects)
             if i not in _routed_project_ids and p.get('routage', 'les_deux') != 'exclure']
if _leftover:
    add_heading(doc, "Exp\u00e9rience pertinente \u2013 Projets de r\u00e9f\u00e9rence", level=1)
    for proj in _leftover:
        _render_project_table(doc, proj)

# Team section
add_heading(doc, "Équipe de projet", level=1)
tbl = doc.add_table(rows=1+len(team_data['equipe']), cols=3)
tbl.style = 'Table Grid'
headers = ['Nom', 'Rôle dans le projet', 'Bureau']
for i, h in enumerate(headers):
    cell = tbl.rows[0].cells[i]
    cell.text = h
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.color.rgb = CHG_BLUE

for i, member in enumerate(team_data['equipe']):
    row = tbl.rows[i+1]
    row.cells[0].text = member.get('nom','')
    row.cells[1].text = member.get('role_dans_projet','')
    _bureau_val = member.get('bureau', '')
    if _bureau_val and 'ind' in _bureau_val.lower():
        _bureau_val = ''
    row.cells[2].text = _bureau_val

doc.add_page_break()

# Annexe A — CVs
add_heading(doc, "Annexe A — Curriculum vitae", level=1)
seen_cvs = set()
for member in team_data['equipe']:
    cv = match_cv(member.get('nom',''), cvs)
    if not cv:
        continue
    cv_nom = cv.get('nom','')
    if cv_nom in seen_cvs:
        continue
    seen_cvs.add(cv_nom)
    add_heading(doc, cv_nom, level=2)
    cv_text = cv.get('text','')
    if cv_text:
        for para in cv_text[:3000].split('\n\n'):
            if para.strip():
                doc.add_paragraph(para.strip())
    doc.add_page_break()

# Conclusion
if content_data.get('conclusion'):
    add_heading(doc, "Mot de la fin", level=1)
    for para in content_data['conclusion'].split('\n\n'):
        if para.strip():
            doc.add_paragraph(para.strip())

# Save
# Session 32 fix: key output_path by FILE_ID (not sanitized filename) to
# eliminate the collision risk fixed for the Go/No-Go report in Session 31 -
# two approvals producing the same sanitized safe_name would otherwise
# silently overwrite each other's OS document.
output_path = f"/local-files/Proposition_CHG_{FILE_ID}.docx"
doc.save(output_path)
print(f"✅ OS proposal saved: {output_path}")

# ── N8N_OUTPUT ────────────────────────────────────────────────────────────────
# Session 51: status/degraded_reasons from sections_ok (STEP 8, line ~555) and
# the AO re-extraction outcome (ao_text, STEP 3) -- both already-existing
# signals, just not previously surfaced. "success" stays the literal string
# for full success, unchanged from today; "partial" is newly introduced on
# paths previously always mislabeled "success". No "error" case here: if
# doc.save() itself throws, the script exits non-zero and pipeline_server.py
# already reports "error" correctly via the subprocess return code.
degraded_reasons = []
if not sections_ok:
    degraded_reasons.append("sections_empty")
if not ao_text:
    degraded_reasons.append("ao_text_unavailable")
if ao_data_parse_failed:
    degraded_reasons.append("ao_data_parse_failed")

result = {
    "status": "success" if (sections_ok and ao_text and not ao_data_parse_failed) else "partial",
    "degraded_reasons": degraded_reasons,
    "ao_client": ao_data.get("client", ""),
    "ao_titre":  ao_data.get("titre_projet", safe_name),
    "output_file": output_path.replace("/local-files/", "/files/"),
    "confirmed_team": [m["nom"] for m in confirmed_team],
    "confirmed_fiches": confirmed_fiches,
}
print(f"N8N_OUTPUT:{json.dumps(result)}")
