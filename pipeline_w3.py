#!/usr/bin/env python3
"""
Session 40: pipeline_w3.py — Addenda regeneration pipeline.

Invoked by pipeline_server.py's _handle_addenda() as a background subprocess:
    python3 pipeline_w3.py <job_id> <original_file_id> <addendum_file_id> <addendum_file_name> <numero_ao>

Scope (per Session 38/39 decisions):
  - Regenerates the Planification document ONLY. Does NOT regenerate the
    Go/No-Go document or the Tableau decisionnel.
  - "After" score is computed by re-running analyze_ao_gonogo() with the
    original AO + addendum combined (the existing multi-file mechanism,
    proven since Session 34) -- NOT by regenerating the Go/No-Go .docx.
  - The regenerated Planification is written to a NEW file keyed by job_id
    (PLANIFICATION_<job_id>.docx), not overwritten in place over the
    original PLANIFICATION_<original_file_id>.docx -- non-destructive by
    default. Confirmed with Pat as a real decision, not silently assumed.

Output: prints "N8N_OUTPUT:{json}" on success, matching the exact
convention pipeline_w1.py/pipeline_server.py already parse from stdout.
Non-zero exit + stderr on failure -- pipeline_server.py's existing
subprocess error handling (in _handle_addenda's run()) already covers this,
no special-casing needed here.
"""
import os, sys, json, io, tempfile, time
import anthropic
from docx import Document
from docx.shared import Pt
from googleapiclient.discovery import build
from google.oauth2 import service_account
from googleapiclient.http import MediaIoBaseDownload

sys.path.insert(0, '/root')
from analyze_ao import analyze_ao_gonogo, extract_ao_text
from planning_agent import run_planning_agent, add_section_title, CHG_BLUE

STATUS_DIR = '/local-files'

if len(sys.argv) < 6:
    print("Usage: pipeline_w3.py <job_id> <original_file_id> <addendum_file_id> <addendum_file_name> <numero_ao>", file=sys.stderr)
    sys.exit(1)

JOB_ID = sys.argv[1]
ORIGINAL_FILE_ID = sys.argv[2]
ADDENDUM_FILE_ID = sys.argv[3]
ADDENDUM_FILE_NAME = sys.argv[4]
NUMERO_AO = sys.argv[5]

print(f"🚀 Addenda pipeline started: job={JOB_ID} original={ORIGINAL_FILE_ID} numero_ao={NUMERO_AO}")

SERVICE_ACCOUNT_FILE = '/root/chg-credentials/service-account.json'
creds = service_account.Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=['https://www.googleapis.com/auth/drive'])
drive = build('drive', 'v3', credentials=creds)
claude = anthropic.Anthropic(max_retries=5)


def download_bytes(file_id):
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, drive.files().get_media(fileId=file_id))
    done = False
    while not done:
        _, done = dl.next_chunk()
    return buf.getvalue()


def save_temp_file(file_bytes, file_name):
    ext = os.path.splitext(file_name)[1].lower()
    if ext not in ('.pdf', '.docx', '.doc', '.xlsx', '.xls'):
        ext = '.pdf'
    tmp = tempfile.NamedTemporaryFile(suffix=ext, delete=False)
    tmp.write(file_bytes)
    tmp.close()
    return tmp.name


def append_addenda_section(doc_path, history, before_score, after_score, latest_summary):
    """Appends an Addenda section to an already-generated Planification .docx,
    reusing planning_agent.py's own heading style (add_section_title, CHG_BLUE)
    so it visually matches rather than looking bolted on. Does not touch
    planning_agent.py itself -- kept fully isolated here per Session 40 design."""
    doc = Document(doc_path)
    add_section_title(doc, "Modifications en lien avec les addendas")

    p_score = doc.add_paragraph()
    p_score.paragraph_format.space_after = Pt(6)
    r1 = p_score.add_run(f"Score Go/No-Go original : {before_score}/100")
    r1.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = CHG_BLUE

    p_change = doc.add_paragraph()
    p_change.paragraph_format.space_after = Pt(8)
    r3 = p_change.add_run(f"Changements apport\u00e9s par ce dernier addendum : {latest_summary}")
    r3.italic = True
    r3.font.size = Pt(9)

    p_hist_title = doc.add_paragraph()
    r4 = p_hist_title.add_run("Historique des addenda :")
    r4.bold = True
    r4.font.size = Pt(9)
    r4.font.color.rgb = CHG_BLUE

    for entry in history:
        p_hist = doc.add_paragraph()
        p_hist.paragraph_format.space_after = Pt(2)
        run = p_hist.add_run(f"Addenda {entry['numero']} : {entry['summary']}")
        run.font.size = Pt(9)

    doc.save(doc_path)


try:
    # ── STEP 1: Load original AO's persisted extraction (before-score source) ──
    ao_data_path = os.path.join(STATUS_DIR, f"AO_DATA_{ORIGINAL_FILE_ID}.json")
    with open(ao_data_path, 'r', encoding='utf-8') as f:
        before_gonogo_data = json.load(f)
    before_score = before_gonogo_data.get('analyse_strategique', {}).get('score_gonogo', 0)
    print(f"✅ Original AO data loaded: before_score={before_score}/100")

    # ── STEP 2: Get original file's name (for extension-correct re-download) ──
    original_paths = []
    if ORIGINAL_FILE_ID.startswith('batch_'):
        batch_manifest_path = os.path.join(STATUS_DIR, f"BATCH_{ORIGINAL_FILE_ID}.json")
        with open(batch_manifest_path, 'r', encoding='utf-8') as f:
            batch_files = json.load(f)
        print(f"AO original multi-fichiers: {len(batch_files)} fichier(s) a re-telecharger")
        for _bf in batch_files:
            print(f"Telechargement AO original: {_bf['file_name']}")
            _bytes = download_bytes(_bf['file_id'])
            original_paths.append(save_temp_file(_bytes, _bf['file_name']))
    else:
        status_path = os.path.join(STATUS_DIR, f"STATUS_{ORIGINAL_FILE_ID}.json")
        with open(status_path, 'r', encoding='utf-8') as f:
            original_status = json.load(f)
        original_file_name = original_status.get('file_name', 'AO_original.pdf')
        print(f"Telechargement AO original: {original_file_name}")
        original_bytes = download_bytes(ORIGINAL_FILE_ID)
        original_paths.append(save_temp_file(original_bytes, original_file_name))

    print(f"Telechargement addendum: {ADDENDUM_FILE_NAME}")
    addendum_bytes = download_bytes(ADDENDUM_FILE_ID)
    addendum_path = save_temp_file(addendum_bytes, ADDENDUM_FILE_NAME)

    print("Recalcul du score avec l'addendum...")
    after_gonogo_data = analyze_ao_gonogo(original_paths + [addendum_path])
    after_score = after_gonogo_data.get('analyse_strategique', {}).get('score_gonogo', 0)
    print(f"✅ Score apres addenda: {after_score}/100 (avant: {before_score}/100)")

    # ── STEP 5: One-line subject summary of the addendum alone ───────────────
    addendum_text = extract_ao_text(addendum_path)
    _summary_prompt = (
        "Resume en une seule phrase courte (maximum 20 mots) l'objet principal "
        "de cet addendum d'appel d'offres, en francais, sans preambule ni "
        "guillemets:\n\n" + addendum_text[:5000]
    )
    _r = claude.messages.create(
        model="claude-sonnet-4-5-20250929", max_tokens=200,
        messages=[{"role": "user", "content": _summary_prompt}]
    )
    addendum_summary = _r.content[0].text.strip()
    print(f"✅ Resume addendum: {addendum_summary}")

    # ── STEP 6: Append to running addenda history ────────────────────────────
    history_path = os.path.join(STATUS_DIR, f"ADDENDA_HISTORY_{ORIGINAL_FILE_ID}.json")
    if os.path.exists(history_path):
        with open(history_path, 'r', encoding='utf-8') as f:
            history = json.load(f)
    else:
        history = []

    new_entry = {
        "numero": len(history) + 1,
        "numero_ao": NUMERO_AO,
        "summary": addendum_summary,
        "addendum_file_name": ADDENDUM_FILE_NAME,
        "before_score": before_score,
        "after_score": after_score,
        "timestamp": time.time()
    }
    history.append(new_entry)

    with open(history_path, 'w', encoding='utf-8') as f:
        json.dump(history, f, ensure_ascii=False)
    print(f"✅ Historique addenda mis a jour: {len(history)} addendum(s) au total")

    # ── STEP 7: Regenerate Planification (based on the "after" data) ─────────
    planning_report_path = os.path.join(STATUS_DIR, f"PLANIFICATION_{JOB_ID}.docx")
    _tmp_json = tempfile.NamedTemporaryFile(suffix='.json', delete=False, mode='w', encoding='utf-8')
    json.dump(after_gonogo_data, _tmp_json, ensure_ascii=False)
    _tmp_json.close()
    run_planning_agent(_tmp_json.name, planning_report_path)
    os.unlink(_tmp_json.name)
    print(f"✅ Planification (base) generee: {planning_report_path}")

    # ── STEP 8: Append Addenda section to the generated document ─────────────
    append_addenda_section(planning_report_path, history, before_score, after_score, addendum_summary)
    print(f"✅ Section Addenda ajoutee: {planning_report_path}")

    # ── N8N_OUTPUT ─────────────────────────────────────────────────────────
    result = {
        "status": "success",
        "numero_ao": NUMERO_AO,
        "original_file_id": ORIGINAL_FILE_ID,
        "addendum_numero": new_entry["numero"],
        "addendum_summary": addendum_summary,
        "before_score": before_score,
        "after_score": after_score,
        "planning_report_file": f"/files/PLANIFICATION_{JOB_ID}.docx"
    }
    print(f"N8N_OUTPUT:{json.dumps(result, ensure_ascii=False)}")

except Exception as e:
    print(f"❌ Addenda pipeline failed: {e}", file=sys.stderr)
    sys.exit(1)
