import json
import os
from datetime import datetime, date, timedelta
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Couleurs CHG ───────────────────────────────────────────────────────────────
CHG_BLUE      = RGBColor(0x1F, 0x49, 0x7D)
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BLUE_BG = "DCE6F1"
CHG_BLUE_BG   = "1F497D"
GREEN_BG      = "E2EFDA"
RED_BG        = "FCE4D6"
ORANGE_BG     = "FFF2CC"
GREY_BG       = "F2F2F2"

# ── Jours fériés QC (pour calcul jours ouvrables) ─────────────────────────────
FERIES_QC = {
    date(2025,12,25), date(2025,12,26),
    date(2026,1,1),   date(2026,1,2),
    date(2026,4,3),   date(2026,5,18),
    date(2026,6,24),  date(2026,7,1),
    date(2026,9,7),   date(2026,10,12),
    date(2026,12,25),
}

def business_days_between(d1, d2):
    """Count business days between two date objects (d1=start, d2=end). QC holidays excluded."""
    count = 0
    current = d1
    while current < d2:
        if current.weekday() < 5 and current not in FERIES_QC:
            count += 1
        current += timedelta(days=1)
    return count

# ── Helpers python-docx ────────────────────────────────────────────────────────

def set_cell_bg(cell, color_hex):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  color_hex)
    tcPr.append(shd)

def set_table_border(table, color='1F497D'):
    tbl   = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),   'single')
        el.set(qn('w:sz'),    '4')
        el.set(qn('w:color'), color)
        tblBorders.append(el)
    tblPr.append(tblBorders)

def set_col_width(table, col_idx, width_cm):
    for row in table.rows:
        row.cells[col_idx].width = Cm(width_cm)

def no_space_para(para):
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after  = Pt(0)

def add_section_title(doc, title):
    p   = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(title)
    run.bold           = True
    run.font.size      = Pt(12)
    run.font.color.rgb = CHG_BLUE
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '6')
    bot.set(qn('w:color'), '1F497D')
    pBdr.append(bot)
    pPr.append(pBdr)

def add_field_row(table, label, value, label_bg=LIGHT_BLUE_BG, use_bullets=False):
    """
    Adds a key-value row to a 2-column table.
    If use_bullets=True and value is a list, renders each item as a bullet.
    If value is a string containing newlines, renders each line as a bullet.
    """
    row = table.add_row()
    set_cell_bg(row.cells[0], label_bg)
    p0  = row.cells[0].paragraphs[0]
    no_space_para(p0)
    run = p0.add_run(label)
    run.bold      = True
    run.font.size = Pt(9)

    cell = row.cells[1]
    if use_bullets and isinstance(value, list) and len(value) > 1:
        for i, item in enumerate(value):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            no_space_para(p)
            r = p.add_run(f"• {item}")
            r.font.size = Pt(9)
    elif isinstance(value, str) and '\n' in value and use_bullets:
        lines = [l.strip() for l in value.split('\n') if l.strip()]
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            no_space_para(p)
            r = p.add_run(f"• {line}")
            r.font.size = Pt(9)
    else:
        p1  = cell.paragraphs[0]
        no_space_para(p1)
        val = value if isinstance(value, str) else (", ".join(value) if isinstance(value, list) else str(value))
        run2 = p1.add_run(val if val else "—")
        run2.font.size = Pt(9)

def cell_para(cell, text, bold=False, size=9, color=None, bg=None, align=None):
    if bg:
        set_cell_bg(cell, bg)
    p   = cell.paragraphs[0]
    no_space_para(p)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text or "—")
    run.bold      = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color

def split_into_segments(text):
    """
    Splits a dense text block into readable segments for bullet rendering.
    Handles comma-separated conditions (e.g. '5 projets = 25 pts, 4 projets = 22 pts')
    and sentence-separated content (e.g. 'Phrase un. Phrase deux.').
    Returns a list of clean strings, or [text] if no meaningful split is found.
    """
    import re
    if not text or not text.strip():
        return []
    text = text.strip()

    # Pattern 1: comma-separated "X = Y" style conditions (common in scoring scales)
    if re.search(r'=\s*[\d.]+\s*pts?', text) and ',' in text:
        segments = [s.strip() for s in text.split(',') if s.strip()]
        if len(segments) > 1:
            return segments

    # Pattern 2: arrow-prefixed lines already present (→) — split on them
    if '→' in text:
        parts = re.split(r'(?=→)', text)
        segments = [p.strip() for p in parts if p.strip()]
        if len(segments) > 1:
            return segments

    # Pattern 3: sentence split (period/semicolon followed by capital letter)
    segments = [s.strip() for s in re.split(r'(?<=[.;])\s+(?=[A-ZÀ-Ü])', text) if s.strip()]
    if len(segments) > 1:
        return segments

    return [text]

def cell_para_bulleted(cell, text, size=9, bg=None, force_split=True):
    """
    Renders text as bullet points inside a cell, splitting dense content
    into readable segments. Falls back to plain paragraph if no split found.
    """
    if bg:
        set_cell_bg(cell, bg)
    segments = split_into_segments(text) if force_split else [text]
    if not segments:
        segments = ["—"]

    if len(segments) == 1:
        p = cell.paragraphs[0]
        no_space_para(p)
        r = p.add_run(segments[0])
        r.font.size = Pt(size)
    else:
        for i, seg in enumerate(segments):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            no_space_para(p)
            r = p.add_run(f"• {seg}")
            r.font.size = Pt(size)

# ── Score qualificatif ─────────────────────────────────────────────────────────

def get_qualificatif(score):
    if score >= 70:
        return "Élevé"
    elif score >= 60:
        return "Moyen"
    else:
        return "Faible"

# ── Score banner ───────────────────────────────────────────────────────────────

def add_score_banner(doc, score, recommandation, nb_incongruites):
    if recommandation == "EXPIRÉ":
        bg, label = "7B7B7B", "EXPIRÉ"
    elif score >= 70:
        bg, label = "375623", "GO"
    elif score >= 50:
        bg, label = "C07000", "CONDITIONNEL"
    else:
        bg, label = "C00000", "NO GO"

    qualificatif = get_qualificatif(score)

    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = 'Table Grid'
    set_table_border(tbl, bg)

    c0, c1, c2 = tbl.rows[0].cells
    for cell in [c0, c1, c2]:
        set_cell_bg(cell, bg)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p0 = c0.paragraphs[0]
    no_space_para(p0)
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(f"SCORE\n{score}/100")
    r0.bold = True; r0.font.size = Pt(16); r0.font.color.rgb = WHITE
    r0b = p0.add_run(f"\n({qualificatif})")
    r0b.bold = False; r0b.font.size = Pt(10); r0b.font.color.rgb = WHITE

    p1 = c1.paragraphs[0]
    no_space_para(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(label)
    r1.bold = True; r1.font.size = Pt(20); r1.font.color.rgb = WHITE

    p2 = c2.paragraphs[0]
    no_space_para(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"{nb_incongruites} incongruité(s)\ndétectée(s)")
    r2.bold = True; r2.font.size = Pt(11); r2.font.color.rgb = WHITE

    set_col_width(tbl, 0, 4)
    set_col_width(tbl, 1, 8)
    set_col_width(tbl, 2, 5)
    doc.add_paragraph()

# ── French date helper ─────────────────────────────────────────────────────────

MOIS_FR = {
    1:"janvier",2:"février",3:"mars",4:"avril",5:"mai",6:"juin",
    7:"juillet",8:"août",9:"septembre",10:"octobre",11:"novembre",12:"décembre"
}

def fmt_date_fr(date_str):
    if not date_str or date_str == "—":
        return date_str or "—"
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"]:
        try:
            d = datetime.strptime(date_str.strip(), fmt)
            return f"{d.day} {MOIS_FR[d.month]} {d.year}"
        except:
            pass
    return date_str

def parse_date_for_compare(date_str):
    if not date_str:
        return None
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"]:
        try:
            return datetime.strptime(date_str.strip(), fmt).date()
        except:
            pass
    return None

# ── Expire banner ──────────────────────────────────────────────────────────────

def add_expire_banner(doc, date_depot, score_analytique, recommandation_analytique):
    RED_BG  = "C00000"
    DARK_BG = "7B0000"
    YELLOW  = RGBColor(0xFF, 0xFF, 0x00)

    tbl = doc.add_table(rows=2, cols=1)
    tbl.style = "Table Grid"
    set_table_border(tbl, RED_BG)

    c0 = tbl.rows[0].cells[0]
    set_cell_bg(c0, RED_BG)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    no_space_para(p0)
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0a = p0.add_run("AO EXPIRÉ")
    r0a.bold = True; r0a.font.size = Pt(18); r0a.font.color.rgb = YELLOW
    date_str = fmt_date_fr(date_depot) if date_depot else "date inconnue"
    r0b = p0.add_run("  —  Date de dépôt passée : " + date_str)
    r0b.bold = False; r0b.font.size = Pt(11); r0b.font.color.rgb = WHITE

    c1 = tbl.rows[1].cells[0]
    set_cell_bg(c1, DARK_BG)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    no_space_para(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(
        "Score analytique (si l'AO avait été actif) :  "
        + str(score_analytique) + "/100  —  " + str(recommandation_analytique)
    )
    r1.bold = True; r1.font.size = Pt(10); r1.font.color.rgb = WHITE
    set_col_width(tbl, 0, 16)
    doc.add_paragraph()

# ── Encadré Projet et services ─────────────────────────────────────────────────

def add_project_summary_box(doc, sommaire_projet, expertises_resume):
    if not sommaire_projet and not expertises_resume:
        return

    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    set_table_border(tbl, "1F497D")
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, "EBF3FB")

    p_title = cell.paragraphs[0]
    no_space_para(p_title)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("PROJET ET SERVICES")
    r_title.bold = True; r_title.font.size = Pt(10); r_title.font.color.rgb = CHG_BLUE

    if sommaire_projet:
        p_som = cell.add_paragraph()
        no_space_para(p_som)
        p_som.paragraph_format.space_before = Pt(4)
        p_som.paragraph_format.space_after  = Pt(6)
        r_label = p_som.add_run("Description du projet : ")
        r_label.bold = True; r_label.font.size = Pt(9)
        r_val = p_som.add_run(sommaire_projet)
        r_val.font.size = Pt(9)

    if expertises_resume:
        p_exp = cell.add_paragraph()
        no_space_para(p_exp)
        p_exp.paragraph_format.space_before = Pt(2)
        r_exp_label = p_exp.add_run("Expertises requises : ")
        r_exp_label.bold = True; r_exp_label.font.size = Pt(9)
        items = expertises_resume if isinstance(expertises_resume, list) else [expertises_resume]
        for expertise in items:
            p_bullet = cell.add_paragraph()
            no_space_para(p_bullet)
            p_bullet.paragraph_format.left_indent = Cm(0.5)
            r_b = p_bullet.add_run(f"• {expertise}")
            r_b.font.size = Pt(9)

    set_col_width(tbl, 0, 16.5)
    doc.add_paragraph()

# ── Tight deadline alert (business days) ──────────────────────────────────────

def add_tight_deadline_alert(doc, date_depot_str):
    """Alert if < 10 business days (QC) between today and deposit date."""
    depot_date = parse_date_for_compare(date_depot_str)
    if not depot_date:
        return
    today = date.today()
    if depot_date <= today:
        return
    bdays = business_days_between(today, depot_date)
    if bdays < 10:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(2)
        r = p.add_run(
            f"⚠  DÉLAI SERRÉ — {bdays} jour(s) ouvrable(s) avant la date de dépôt. "
            "Mobilisation immédiate requise."
        )
        r.bold = True; r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

# ── Main report generator ──────────────────────────────────────────────────────

def generate_gonogo_report(json_path, output_path):
    with open(json_path, encoding="utf-8") as f:
        data = json.load(f)

    ident   = data.get("identification", {})
    desc    = data.get("description_projet", {})
    conf    = data.get("criteres_conformite", {})
    eval_   = data.get("criteres_evaluation", {})
    ress    = data.get("ressources_requises", {})
    incon   = data.get("incongruites", [])
    strat   = data.get("analyse_strategique", {})
    ech     = data.get("echeancier_redaction", {})
    risques = data.get("risques_contractuels", {})

    score          = strat.get("score_gonogo", 0)
    recommandation = strat.get("recommandation", "")
    nb_incon       = len(incon)
    date_depot_raw = ident.get("date_depot", "")

    doc = Document()

    # ── Marges ────────────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── En-tête avec logo ─────────────────────────────────────────────────────
    logo_path = "/root/chg-library/Logo_CHG.jpg"
    header    = doc.sections[0].header
    htbl      = header.add_table(1, 2, Cm(16))
    htbl.style = 'Table Grid'
    set_table_border(htbl, 'FFFFFF')

    hc0 = htbl.rows[0].cells[0]
    if os.path.exists(logo_path):
        hp = hc0.paragraphs[0]
        run = hp.add_run()
        run.add_picture(logo_path, width=Cm(4))
    else:
        hc0.paragraphs[0].add_run("Groupe Conseil CHG").bold = True

    hc1 = htbl.rows[0].cells[1]
    hp1 = hc1.paragraphs[0]
    hp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r  = hp1.add_run("ANALYSE GO/NO-GO")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = CHG_BLUE
    now = datetime.now()
    hp1.add_run(f"\nGénéré le {now.day} {MOIS_FR[now.month]} {now.year}").font.size = Pt(8)

    # ── Titre principal (Numéro AO, Titre, Client restent ici) ───────────────
    titre = doc.add_paragraph()
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = titre.add_run(ident.get("titre_projet", "Analyse Go/No-Go").upper())
    tr.bold = True; tr.font.size = Pt(14); tr.font.color.rgb = CHG_BLUE
    titre.paragraph_format.space_before = Pt(6)
    titre.paragraph_format.space_after  = Pt(2)

    sous_titre = doc.add_paragraph()
    sous_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st = sous_titre.add_run(
        f"{ident.get('client','')}  |  AO {ident.get('numero_ao','')}"
    )
    st.font.size = Pt(10); st.font.color.rgb = CHG_BLUE
    sous_titre.paragraph_format.space_after = Pt(8)

    # ── Bannières ─────────────────────────────────────────────────────────────
    if recommandation == "EXPIRÉ":
        score_analytique = strat.get("score_analytique", score)
        reco_analytique  = strat.get("recommandation_analytique", "—")
        add_expire_banner(doc, date_depot_raw, score_analytique, reco_analytique)

    add_score_banner(doc, score, recommandation, nb_incon)

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 1 — Appel d'offres et exigences administratives
    # Contains: AO info + conformité + pénalités (all in one table)
    # Does NOT contain: Numéro AO, Titre projet, Client (those are in title)
    # ══════════════════════════════════════════════════════════════════════════
    add_section_title(doc, "1. APPEL D'OFFRES ET EXIGENCES ADMINISTRATIVES")

    tbl1 = doc.add_table(rows=0, cols=2)
    tbl1.style = 'Table Grid'
    set_table_border(tbl1)
    set_col_width(tbl1, 0, 5.5)
    set_col_width(tbl1, 1, 11)

    # AO info
    add_field_row(tbl1, "Contact",
                  f"{ident.get('contact_client','')}  {ident.get('email_contact','')}  {ident.get('telephone_contact','')}".strip())
    add_field_row(tbl1, "Date de dépôt",
                  (fmt_date_fr(ident.get("date_depot","")) + f"  {ident.get('heure_depot','')}").strip())
    add_field_row(tbl1, "Format de dépôt", ident.get("format_depot",""))
    add_field_row(tbl1, "Lieu de dépôt",   ident.get("lieu_depot",""))

    if ident.get("soumission_physique") or ident.get("delai_purolator_requis"):
        add_field_row(tbl1, "⚠ Livraison des documents (Purolator)",
                      "+2 jours ouvrables requis pour livraison physique",
                      label_bg="FCE4D6")

    tz = ident.get("timezone_projet","")
    if tz and tz.upper() not in ["", "EST", "ET", "EASTERN"]:
        add_field_row(tbl1, "⚠ FUSEAU HORAIRE",
                      f"Attention: {tz} — vérifier l'heure de dépôt",
                      label_bg="FCE4D6")

    add_field_row(tbl1, "Date limite questions", fmt_date_fr(ident.get("date_limite_questions","")))
    add_field_row(tbl1, "Date d'émission",       fmt_date_fr(ident.get("date_emission","")))
    add_field_row(tbl1, "Localisation projet",   desc.get("localisation",""))
    add_field_row(tbl1, "Région admin.",          desc.get("region_administrative",""))
    add_field_row(tbl1, "Budget estimé",          desc.get("budget_estime",""))

    # Conformité (merged into Table 1)
    add_field_row(tbl1, "Visite obligatoire",
                  ("Oui — " + fmt_date_fr(conf.get("date_visite",""))) if conf.get("visite_obligatoire") else "Non")
    add_field_row(tbl1, "Garantie de soumission",
                  f"Oui — {conf.get('montant_garantie','')}" if conf.get("garantie_soumission_requise") else "Non")

    assurances = conf.get("assurances_requises", [])
    montant_rc = conf.get("montant_assurance_responsabilite","")
    assur_list = list(assurances) if assurances else []
    if montant_rc:
        assur_list.append(f"RC Professionnelle: {montant_rc}")
    bg_assur = "FCE4D6" if assurances else LIGHT_BLUE_BG
    add_field_row(tbl1, "Assurances requises",
                  assur_list if assur_list else ["Non spécifié"],
                  label_bg=bg_assur, use_bullets=True)

    certs = conf.get("formations_certifications_requises", [])
    add_field_row(tbl1, "Certifications requises",
                  certs if certs else ["Aucune"],
                  use_bullets=True)

    elim = conf.get("criteres_eliminatoires", [])
    if elim:
        add_field_row(tbl1, "⚠ CRITÈRES ÉLIMINATOIRES",
                      elim, label_bg="FCE4D6", use_bullets=True)

    # Pénalités (from risques_contractuels)
    penalites = risques.get("penalites","") or "Non mentionnées"
    add_field_row(tbl1, "Pénalités", penalites)

    # ── Encadré Projet et services (sous Section 1) ───────────────────────────
    doc.add_paragraph()
    sommaire   = desc.get("sommaire_projet","")
    expertises = desc.get("expertises_requises_resume", [])
    add_project_summary_box(doc, sommaire, expertises)

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 2 — Critères d'évaluation
    # ══════════════════════════════════════════════════════════════════════════
    add_section_title(doc, "2. CRITÈRES D'ÉVALUATION")

    tbl2 = doc.add_table(rows=0, cols=2)
    tbl2.style = 'Table Grid'
    set_table_border(tbl2)
    set_col_width(tbl2, 0, 5.5)
    set_col_width(tbl2, 1, 11)
    add_field_row(tbl2, "Mode d'attribution",     eval_.get("mode_attribution",""))
    add_field_row(tbl2, "Pondération technique",  eval_.get("ponderation_technique",""))
    add_field_row(tbl2, "Pondération financière", eval_.get("ponderation_financiere",""))
    add_field_row(tbl2, "Note minimale qualité",  eval_.get("note_minimale_qualite",""))
    add_field_row(tbl2, "Entrevue prévue",        "Oui" if eval_.get("entrevue_prevue") else "Non")
    add_field_row(tbl2, "Organigramme requis",    "Oui" if eval_.get("organigramme_requis") else "Non")

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 3 — Ressources et expertises requises
    # ══════════════════════════════════════════════════════════════════════════
    add_section_title(doc, "3. RESSOURCES ET EXPERTISES REQUISES")

    tbl3a = doc.add_table(rows=0, cols=2)
    tbl3a.style = 'Table Grid'
    set_table_border(tbl3a)
    set_col_width(tbl3a, 0, 5.5)
    set_col_width(tbl3a, 1, 11)

    add_field_row(tbl3a, "CV requis",
                  f"Oui{'  — avec organigramme' if eval_.get('organigramme_requis') else ''}"
                  if ress.get("cv_requis") else "Non")
    add_field_row(tbl3a, "Nombre de CV", ress.get("nombre_cv","Non spécifié"))
    add_field_row(tbl3a, "Fiches projets requises",
                  f"Oui{'  — ' + ress.get('montant_minimal_projets','') if ress.get('montant_minimal_projets') else ''}"
                  if ress.get("fiches_projets_requises") else "Non")
    add_field_row(tbl3a, "Sous-traitance autorisée",
                  "Oui" if ress.get("sous_traitance_autorisee") else "Non")

    if ress.get("sous_traitance_recommandee"):
        domaines = ress.get("domaines_sous_traitance", [])
        add_field_row(tbl3a, "⚠ SOUS-TRAITANCE RECOMMANDÉE",
                      "Domaines: " + (", ".join(domaines) if domaines else "voir analyse"),
                      label_bg="FFF2CC")

    # Consortium — always shown
    add_field_row(tbl3a, "Consortium autorisé",
                  "Oui" if risques.get("consortium_autorise") else "Non")

    # Rôles détaillés
    roles = ress.get("roles_detailles", [])
    if roles:
        doc.add_paragraph()
        p_roles = doc.add_paragraph()
        no_space_para(p_roles)
        r_roles = p_roles.add_run("Rôles requis par l'appel d'offres :")
        r_roles.bold = True; r_roles.font.size = Pt(10); r_roles.font.color.rgb = CHG_BLUE

        tbl3b = doc.add_table(rows=1, cols=5)
        tbl3b.style = 'Table Grid'
        set_table_border(tbl3b)
        hdrs = ["Rôle", "Titre requis", "Exp. min.", "Qualifications obligatoires", "Points / Différenciateurs"]
        for i, h in enumerate(hdrs):
            c = tbl3b.rows[0].cells[i]
            set_cell_bg(c, CHG_BLUE_BG)
            cell_para(c, h, bold=True, color=WHITE, align='center')
        set_col_width(tbl3b, 0, 3)
        set_col_width(tbl3b, 1, 2.5)
        set_col_width(tbl3b, 2, 1.5)
        set_col_width(tbl3b, 3, 5)
        set_col_width(tbl3b, 4, 5)

        for role in roles:
            row = tbl3b.add_row()
            cell_para(row.cells[0], role.get("role",""))
            cell_para(row.cells[1], role.get("titre_requis",""))
            ans = role.get("annees_experience_min", 0)
            ans_txt = (f"{ans} ans" if role.get("annees_experience_mentionnees") and ans
                       else "Non spécifié")
            cell_para(row.cells[2], ans_txt, align='center')
            quals = role.get("qualifications_obligatoires", [])
            # Bullet points for qualifications
            if quals and len(quals) > 1:
                c = row.cells[3]
                set_cell_bg(c, "FFFFFF")
                for i, q in enumerate(quals):
                    p = c.paragraphs[0] if i == 0 else c.add_paragraph()
                    no_space_para(p)
                    p.add_run(f"• {q}").font.size = Pt(9)
            else:
                cell_para(row.cells[3], quals[0] if quals else "Aucune spécifiée")

            pts_diff = role.get("points_differenciateurs","")
            bg_pts = ORANGE_BG if pts_diff and pts_diff.strip() else LIGHT_BLUE_BG
            cell_para_bulleted(row.cells[4], pts_diff if pts_diff else "—", bg=bg_pts)

        # Note under roles table
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(2)
        r_note = p_note.add_run(
            "* Colonne Points / Différenciateurs : liens automatiques avec les critères "
            "d'évaluation — en cours d'optimisation."
        )
        r_note.italic = True; r_note.font.size = Pt(8)
        r_note.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 4 — Détail des critères d'évaluation
    # ══════════════════════════════════════════════════════════════════════════
    add_section_title(doc, "4. DÉTAIL DES CRITÈRES D'ÉVALUATION")

    criteres = eval_.get("criteres_detail", [])
    if criteres:
        tbl4 = doc.add_table(rows=1, cols=4)
        tbl4.style = 'Table Grid'
        set_table_border(tbl4)
        hdrs = ["#", "Critère", "Points", "Contenu requis / Exigences"]
        for i, h in enumerate(hdrs):
            c = tbl4.rows[0].cells[i]
            set_cell_bg(c, CHG_BLUE_BG)
            cell_para(c, h, bold=True, color=WHITE, align='center')
        set_col_width(tbl4, 0, 1)
        set_col_width(tbl4, 1, 4)
        set_col_width(tbl4, 2, 1.5)
        set_col_width(tbl4, 3, 10)

        for crit in criteres:
            row = tbl4.add_row()
            cell_para(row.cells[0], str(crit.get("numero_critere","")), align='center')
            cell_para(row.cells[1], crit.get("intitule",""))
            cell_para(row.cells[2], str(crit.get("points_max","")), align='center')
            details = crit.get("contenu_requis","")
            nb_proj = crit.get("nb_projets_requis", 0)
            nb_ans  = crit.get("annees_experience_requises", 0)
            if nb_proj:
                details += f"\n→ {nb_proj} projet(s) comparable(s) requis"
            if nb_ans:
                details += f"\n→ {nb_ans} an(s) d'expérience requis"
            cell_para_bulleted(row.cells[3], details)
    else:
        p = doc.add_paragraph()
        r = p.add_run(
            "Aucun critère détecté — relancer l'analyse avec un AO complet."
        )
        r.italic = True; r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 5 — Échéancier de rédaction
    # Order: Réunion → Retour ingénieur → Document final → ISO → Purolator → DÉPÔT
    # Document final = Retour ingénieur + 1 business day (calculated in analyze_ao.py)
    # ══════════════════════════════════════════════════════════════════════════
    add_section_title(doc, "5. ÉCHÉANCIER DE RÉDACTION")

    today_str = date.today().strftime("%Y-%m-%d")

    if ident.get("soumission_physique") or ident.get("delai_purolator_requis"):
        p_pur = doc.add_paragraph()
        r_pur = p_pur.add_run(
            "⚠ Soumission physique — délai Purolator (+2 jours ouvrables) déjà intégré dans l'échéancier"
        )
        r_pur.bold = True; r_pur.font.size = Pt(9)
        r_pur.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    tbl5 = doc.add_table(rows=1, cols=3)
    tbl5.style = 'Table Grid'
    set_table_border(tbl5)
    for i, h in enumerate(["Étape", "Date cible", "Statut"]):
        c = tbl5.rows[0].cells[i]
        set_cell_bg(c, CHG_BLUE_BG)
        cell_para(c, h, bold=True, color=WHITE, align='center')
    set_col_width(tbl5, 0, 8)
    set_col_width(tbl5, 1, 4)
    set_col_width(tbl5, 2, 4.5)

    # Confirmed order from Alexandra
    is_physical = ident.get("soumission_physique") or ident.get("delai_purolator_requis")
    etapes = [
        ("Réunion de planification",     ech.get("date_reunion_planification","")),
        ("Retour ingénieur responsable", ech.get("date_limite_ingenieur","")),
        ("Document final",               ech.get("date_document_final","")),
        ("Vérification ISO / qualité",   ech.get("date_iso","")),
        ("Livraison Purolator",          ech.get("date_livraison_purolator","") if is_physical else None),
        ("DATE DE DÉPÔT",                ech.get("date_depot","")),
    ]

    for label, date_val in etapes:
        if date_val is None:
            continue
        row = tbl5.add_row()
        is_depot = "DÉPÔT" in label

        if is_depot:
            set_cell_bg(row.cells[0], CHG_BLUE_BG)
            cell_para(row.cells[0], label, bold=True, color=WHITE)
        else:
            cell_para(row.cells[0], label)

        date_display = fmt_date_fr(date_val) if date_val else "À calculer"
        cell_para(row.cells[1], date_display, bold=is_depot, align='center')

        statut, bg_statut = "", LIGHT_BLUE_BG
        if date_val:
            d = parse_date_for_compare(date_val)
            if d:
                statut     = "PASSÉ"  if d.strftime("%Y-%m-%d") < today_str else "À FAIRE"
                bg_statut  = RED_BG   if statut == "PASSÉ"                  else GREEN_BG
            else:
                statut = "À VÉRIFIER"
        else:
            statut = "À CALCULER"

        set_cell_bg(row.cells[2], bg_statut)
        cell_para(row.cells[2], statut, bold=True, align='center')

    notes = ech.get("notes_echeancier","")
    if notes:
        p_n = doc.add_paragraph()
        r_n = p_n.add_run(f"Note : {notes}")
        r_n.italic = True; r_n.font.size = Pt(8)

    # Tight deadline alert (< 10 business days)
    add_tight_deadline_alert(doc, date_depot_raw)

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 6 — Analyse stratégique
    # ══════════════════════════════════════════════════════════════════════════
    add_section_title(doc, "6. ANALYSE STRATÉGIQUE")

    forts   = strat.get("points_forts_chg", [])
    faibles = strat.get("points_faibles_chg", [])
    max_rows = max(len(forts), len(faibles), 1)

    tbl6 = doc.add_table(rows=1, cols=2)
    tbl6.style = "Table Grid"
    set_table_border(tbl6)
    set_cell_bg(tbl6.rows[0].cells[0], "375623")
    set_cell_bg(tbl6.rows[0].cells[1], "C00000")
    cell_para(tbl6.rows[0].cells[0], "POINTS FORTS CHG",    bold=True, color=WHITE, align="center")
    cell_para(tbl6.rows[0].cells[1], "POINTS DE VIGILANCE", bold=True, color=WHITE, align="center")
    set_col_width(tbl6, 0, 8.5)
    set_col_width(tbl6, 1, 8.5)

    for i in range(max_rows):
        row = tbl6.add_row()
        set_cell_bg(row.cells[0], GREEN_BG)
        set_cell_bg(row.cells[1], RED_BG)
        cell_para(row.cells[0], f"• {forts[i]}"   if i < len(forts)   else "")
        cell_para(row.cells[1], f"• {faibles[i]}" if i < len(faibles) else "")

    # Concurrence placeholder (Phase 2)
    concurrents = strat.get("concurrents_region", [])
    doc.add_paragraph()
    p_conc_title = doc.add_paragraph()
    no_space_para(p_conc_title)
    r_conc = p_conc_title.add_run("Analyse de la concurrence :")
    r_conc.bold = True; r_conc.font.size = Pt(10); r_conc.font.color.rgb = CHG_BLUE

    if concurrents:
        tbl_conc = doc.add_table(rows=1, cols=3)
        tbl_conc.style = "Table Grid"
        set_table_border(tbl_conc)
        for i, h in enumerate(["Concurrent", "Région", "Historique CHG"]):
            c = tbl_conc.rows[0].cells[i]
            set_cell_bg(c, CHG_BLUE_BG)
            cell_para(c, h, bold=True, color=WHITE, align='center')
        set_col_width(tbl_conc, 0, 5)
        set_col_width(tbl_conc, 1, 4)
        set_col_width(tbl_conc, 2, 8)
        for comp in concurrents:
            row = tbl_conc.add_row()
            cell_para(row.cells[0], comp.get("nom",""))
            cell_para(row.cells[1], comp.get("region",""))
            cell_para(row.cells[2], comp.get("historique",""))
    else:
        p_conc2 = doc.add_paragraph()
        p_conc2.paragraph_format.space_before = Pt(2)
        r_conc2 = p_conc2.add_run(
            "Données de concurrence à intégrer — Phase 2 "
            "(source : LISTE GÉNÉRALE OS ET PROJETS.xlsx)"
        )
        r_conc2.italic = True; r_conc2.font.size = Pt(9)
        r_conc2.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 7 — Incongruités et questions
    # ══════════════════════════════════════════════════════════════════════════
    add_section_title(doc, f"7. INCONGRUITÉS ET QUESTIONS ({nb_incon} détectée(s))")

    if not incon:
        p = doc.add_paragraph()
        p.add_run("Aucune incongruité majeure détectée.").italic = True
    else:
        tbl7 = doc.add_table(rows=1, cols=3)
        tbl7.style = 'Table Grid'
        set_table_border(tbl7)
        for i, h in enumerate(["#", "Problème détecté", "Question à poser au client"]):
            c = tbl7.rows[0].cells[i]
            set_cell_bg(c, CHG_BLUE_BG)
            cell_para(c, h, bold=True, color=WHITE, align='center')
        set_col_width(tbl7, 0, 1)
        set_col_width(tbl7, 1, 8)
        set_col_width(tbl7, 2, 7.5)

        for idx, item in enumerate(incon, 1):
            row = tbl7.add_row()
            set_cell_bg(row.cells[0], LIGHT_BLUE_BG)
            cell_para(row.cells[0], str(idx), bold=True, align='center')
            section_txt = item.get("section","")
            desc_txt    = item.get("description","")
            full_txt    = f"[{section_txt}]\n{desc_txt}" if section_txt else desc_txt
            cell_para(row.cells[1], full_txt)
            cell_para(row.cells[2], item.get("question_a_poser",""))

    # ── Confidentialité (hardcodé) ────────────────────────────────────────────
    doc.add_paragraph()
    p_conf = doc.add_paragraph()
    p_conf.paragraph_format.space_before = Pt(12)
    r_conf = p_conf.add_run(
        "CONFIDENTIALITÉ — Cette analyse est la propriété du Groupe Conseil CHG inc. "
        "Elle contient des informations stratégiques confidentielles et est destinée "
        "exclusivement aux personnes autorisées de la firme. "
        "Toute reproduction ou diffusion externe est strictement interdite."
    )
    r_conf.italic = True; r_conf.font.size = Pt(8)
    r_conf.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(output_path)
    print(f"Rapport Go/No-Go généré: {output_path}")
    return output_path


if __name__ == "__main__":
    import sys
    json_path   = sys.argv[1] if len(sys.argv) > 1 else "/root/test_ao_gonogo.json"
    output_path = sys.argv[2] if len(sys.argv) > 2 else "/local-files/LATEST_GONOGO_RAPPORT.docx"
    generate_gonogo_report(json_path, output_path)
