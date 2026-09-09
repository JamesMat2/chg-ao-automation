import os, sys, json, io, tempfile, datetime, time, subprocess
import anthropic
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from googleapiclient.discovery import build
from google.oauth2 import service_account
from googleapiclient.http import MediaIoBaseDownload
sys.path.insert(0, '/root')
from analyze_ao import (
    analyze_ao_gonogo, TruncatedResponseError, ResponseRefusedError,
    apply_refusal_substitutions, restore_refusal_substitutions,
)
from gonogo_report import generate_gonogo_report
from planning_agent import run_planning_agent

# Args from n8n
FILE_ID = sys.argv[1] if len(sys.argv) > 1 else None
FILE_NAME = sys.argv[2] if len(sys.argv) > 2 else "AO_inconnu"
BATCH_FILES_JSON = sys.argv[3] if len(sys.argv) > 3 else None

print(f"🚀 Pipeline v2 started for: {FILE_NAME} ({FILE_ID})")
if BATCH_FILES_JSON:
    print(f"📦 Mode multi-fichiers actif — liste: {BATCH_FILES_JSON}")

SERVICE_ACCOUNT_FILE = '/root/chg-credentials/service-account.json'
creds = service_account.Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=['https://www.googleapis.com/auth/drive'])
drive = build('drive', 'v3', credentials=creds)
claude = anthropic.Anthropic(max_retries=5)

cvs_raw = json.load(open('/root/chg-library/cv_docx_library.json'))
cvs = [{'nom': cv.get('nom', ''), 'titre': cv.get('nom', ''), 'competences': [], 'text': cv.get('texte', '')} for cv in cvs_raw]
projects = json.load(open('/root/chg-library/projects_library.json'))
real_proposal = json.load(open('/root/chg-library/real_proposal_analysis.json'))
crossref = json.load(open('/root/chg-library/ao_os_crossref.json'))
crossref_intel = {'lecons': [], 'arguments_differenciateurs': [], 'mots_cles_techniques': [], 'mappings': []}
for pair in crossref.get('pairs', []):
    crossref_intel['lecons'].extend(pair.get('lecons_pour_futurs_aos', []))
    crossref_intel['arguments_differenciateurs'].extend(pair.get('arguments_differenciateurs', []))
    crossref_intel['mots_cles_techniques'].extend(pair.get('mots_cles_techniques', []))
    crossref_intel['mappings'].extend(pair.get('comment_chg_repond_aux_exigences', []))
crossref_intel['lecons'] = list(dict.fromkeys(crossref_intel['lecons']))[:20]
crossref_intel['arguments_differenciateurs'] = list(dict.fromkeys(crossref_intel['arguments_differenciateurs']))[:10]
crossref_intel['mots_cles_techniques'] = list(dict.fromkeys(crossref_intel['mots_cles_techniques']))[:15]

def download_bytes(file_id):
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, drive.files().get_media(fileId=file_id))
    done = False
    while not done: _, done = dl.next_chunk()
    return buf.getvalue()

def pdf_to_text(pdf_bytes, max_chars=10000):
    try:
        from pdfminer.high_level import extract_text as extract_text_pdf
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
            f.write(pdf_bytes); tmp = f.name
        text = extract_text_pdf(tmp); os.unlink(tmp)
        return text[:max_chars]
    except Exception as e:
        print(f"Erreur extraction PDF: {e}")
        return ""

def docx_to_text(docx_bytes, max_chars=10000):
    import io as _io
    from docx import Document as _DocxDoc
    try:
        doc = _DocxDoc(_io.BytesIO(docx_bytes))
        text = chr(10).join([p.text for p in doc.paragraphs if p.text.strip()])
        return text[:max_chars]
    except Exception as e:
        print(f"Warning: docx extraction error: {e}")
        return ""

def claude_call(prompt, max_tokens=2000):
    r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}])
    resp = r.content[0].text.strip()
    if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
    return resp

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        run.font.bold = True

def add_divider(doc):
    p = doc.add_paragraph('—' * 80)
    p.runs[0].font.size = Pt(6)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

def match_cv(member_name, cvs):
    member_lower = member_name.lower()
    member_base = member_lower.split(' - ')[0].strip()
    member_words = set(member_base.split())
    best_cv = None
    best_score = 0
    for c in cvs:
        cv_base = c['nom'].lower().split(' - ')[0].strip()
        cv_words = set(cv_base.split())
        score = len(member_words & cv_words)
        if score > best_score:
            best_score = score
            best_cv = c
    return best_cv if best_score >= 1 else None

# ── STEP 1: Download and extract AO text ──────────────────────────────────────
def docx_to_text_inner(docx_bytes, max_chars=10000):
    import io as _io
    from docx import Document as _D
    try:
        doc = _D(_io.BytesIO(docx_bytes))
        return chr(10).join([p.text for p in doc.paragraphs if p.text.strip()])[:max_chars]
    except Exception as e:
        print(f"Erreur extraction DOCX: {e}")
        return ""

def _save_temp_file(file_bytes, file_name):
    _ext = os.path.splitext(file_name)[1].lower()
    if _ext not in ('.pdf', '.docx', '.doc', '.xlsx', '.xls'):
        _ext = '.pdf'
    _tmp = tempfile.NamedTemporaryFile(suffix=_ext, delete=False)
    _tmp.write(file_bytes); _tmp.close()
    return _tmp.name

if BATCH_FILES_JSON:
    with open(BATCH_FILES_JSON, 'r', encoding='utf-8') as _bf:
        _batch_files = json.load(_bf)
    print(f"📦 {len(_batch_files)} fichier(s) dans ce batch")
    ao_pdf_paths = []
    _main_bytes = None
    _main_name = None
    for _i, _bfitem in enumerate(_batch_files):
        _fid = _bfitem['file_id']; _fname = _bfitem['file_name']
        print(f"📄 Téléchargement ({_i+1}/{len(_batch_files)}): {_fname}")
        _fbytes = download_bytes(_fid)
        ao_pdf_paths.append(_save_temp_file(_fbytes, _fname))
        if _i == 0:
            _main_bytes, _main_name = _fbytes, _fname
    ao_pdf_path = ao_pdf_paths[0]
    if _main_name.lower().endswith(('.docx', '.doc')):
        ao_text = docx_to_text_inner(_main_bytes)
    else:
        ao_text = pdf_to_text(_main_bytes)[:10000]
    print(f"📄 Document principal du batch: {_main_name}")
else:
    print("📄 Downloading AO...")
    ao_bytes = download_bytes(FILE_ID)
    is_docx = FILE_NAME.lower().endswith('.docx') or FILE_NAME.lower().endswith('.doc')
    if is_docx:
        ao_pdf_path = _save_temp_file(ao_bytes, FILE_NAME)
        ao_text = docx_to_text_inner(ao_bytes)
        print(f"📝 Detected Word document — extracting text from .docx")
    else:
        ao_pdf_path = _save_temp_file(ao_bytes, FILE_NAME)
        ao_text = pdf_to_text(ao_bytes)[:10000]
        print(f"📄 Detected PDF — extracting text with pdfminer")
    ao_pdf_paths = [ao_pdf_path]

# Session 51: pdf_to_text()/docx_to_text_inner() return "" on any extraction
# failure rather than raising (Part A above now at least logs it). Left
# unguarded, an empty ao_text still flows into STEP 2's Claude call, which
# reliably returns valid-but-blank JSON on empty input rather than raising --
# so nothing downstream would ever surface this as an error. Verified via
# pipeline_server.py's exit-code check (catches a bare, uncaught raise here
# unconditionally, independent of whether any N8N_OUTPUT line was printed)
# and n8n's 25-minute-capped poll loop (Notify Failure email either way) --
# no wrapping try/except needed around STEP 1/2 for this to fail cleanly.
if not ao_text:
    raise Exception(f"Impossible d'extraire le texte du document principal: {_main_name if BATCH_FILES_JSON else FILE_NAME}")

# ── STEP 2: Analyze AO ────────────────────────────────────────────────────────
print("🧠 Analyzing AO with Claude...")
_ao_call_prompt = f"AO:\n{ao_text}\n\nUNIQUEMENT JSON sans backticks:\n{{\"client\":\"\",\"numero_ao\":\"\",\"titre_projet\":\"\",\"date_limite\":\"\",\"type_travaux\":\"\",\"competences_requises\":[],\"localisation\":\"\",\"resume_mandat\":\"\",\"points_addendas\":[]}}"

def _step2_call(prompt_text):
    _r = None
    for _attempt, _wait in enumerate([0, 5, 15], start=1):
        if _wait:
            print(f"  [retry] Claude API issue — waiting {_wait}s before attempt {_attempt}/3")
            time.sleep(_wait)
        try:
            _r = claude.messages.create(model="claude-sonnet-4-5-20250929", max_tokens=2000,
                messages=[{"role": "user", "content": prompt_text}])
            break
        except anthropic.APIStatusError as _api_e:
            if _attempt == 3:
                raise
            print(f"  [retry] Claude API call failed ({_api_e}) — will retry")
    # Session 52: confirmed via a real production incident (AO 13203) that
    # Claude can return stop_reason == "refusal" with zero content blocks --
    # not an APIStatusError, so the retry loop above never sees it (the SDK
    # treats a refusal as a normal completed response). Unguarded, this
    # crashed with a bare IndexError on r.content[0] below.
    if not _r.content:
        if _r.stop_reason == "refusal":
            raise ResponseRefusedError(
                f"STEP 2: Claude refused to respond (stop_reason=refusal) — "
                f"content policy trigger, not a parse/truncation issue."
            )
        raise ResponseRefusedError(f"STEP 2: Claude response has no content (stop_reason={_r.stop_reason!r})")
    return _r

try:
    r = _step2_call(_ao_call_prompt)
except ResponseRefusedError as _re:
    print(f"  [refusal retry] STEP 2 initial call refused ({_re}) — retrying once with known-trigger-phrase substitution")
    r = _step2_call(apply_refusal_substitutions(_ao_call_prompt))

resp = r.content[0].text.strip()
if resp.startswith("```"): resp = resp.split("\n",1)[1].rsplit("```",1)[0].strip()
# Session 52: this parse/access step has no retry of its own (only the API
# call above does) -- any failure here already propagates uncaught exactly
# as before. The only change: if the response was confirmed truncated by
# the API itself (stop_reason == "max_tokens"), the resulting crash is
# re-raised with that diagnosis attached instead of a generic parse error --
# no retry-loop or token-budget change (Option A, scoped deliberately
# narrower than the analyze_ao.py fix -- see Session 52 scoping notes).
try:
    ao_data = json.loads(resp)
    # Session 52: restore any refusal-trigger placeholder back to the real
    # phrase immediately -- this is pipeline_w1.py's only consumer of
    # ao_data (the debug print and N8N_OUTPUT's ao_client/ao_titre below),
    # so restoring right here is the single boundary point for this script.
    ao_data = restore_refusal_substitutions(ao_data)
    print(f"DEBUG ao_data keys: {list(ao_data.keys())}")
    print(f"✅ AO: {ao_data['client']} — {ao_data['titre_projet']}")
except (json.JSONDecodeError, KeyError) as _parse_e:
    if r.stop_reason == "max_tokens":
        raise TruncatedResponseError(f"STEP 2 response truncated at max_tokens=2000 ({_parse_e})") from _parse_e
    raise

# ── STEP 3: Go/No-Go Analysis ─────────────────────────────────────────────────
print("\n🔍 Running Go/No-Go analysis...")
gonogo_score = 0; gonogo_recommandation = ""; gonogo_nb_incongruites = 0; gonogo_report_path = None; planning_report_path = None
try:
    gonogo_data = analyze_ao_gonogo(ao_pdf_paths)
    # Session 39: persist full extraction for later addenda regeneration
    try:
        with open(f"/local-files/AO_DATA_{FILE_ID}.json", "w", encoding="utf-8") as _aodata_f:
            json.dump(gonogo_data, _aodata_f, ensure_ascii=False)
        print(f"\u2705 AO data persisted: /local-files/AO_DATA_{FILE_ID}.json")
    except Exception as _aodata_e:
        print(f"\u26a0\ufe0f AO_DATA persist failed (non-blocking): {_aodata_e}")
    # Session 39 fix: score_gonogo and recommandation are nested inside
    # analyse_strategique, not top-level -- previously silently defaulted to 0/''
    gonogo_score = gonogo_data.get('analyse_strategique', {}).get('score_gonogo', 0)
    gonogo_recommandation = gonogo_data.get('analyse_strategique', {}).get('recommandation', '')
    gonogo_nb_incongruites = len(gonogo_data.get('incongruites', []))
    print(f"✅ Go/No-Go: {gonogo_score}/100 — {gonogo_recommandation}")
    gonogo_report_path = f"/local-files/GONOGO_RAPPORT_{FILE_ID}.docx"
    import tempfile as _tf
    _tmp_json = _tf.NamedTemporaryFile(suffix='.json', delete=False, mode='w')
    json.dump(gonogo_data, _tmp_json, ensure_ascii=False)
    _tmp_json.close()
    generate_gonogo_report(_tmp_json.name, gonogo_report_path)
    os.unlink(_tmp_json.name)
    print(f"✅ Go/No-Go report: {gonogo_report_path}")


except Exception as e:
    print(f"⚠️ Go/No-Go failed (non-blocking): {e}")
    gonogo_data = {}
    gonogo_report_path = None
# ── Agent 3 — Tableau décisionnel ────────────────────────────────────────────
tableau_report_path = None
try:
    _tableau_template = "/root/chg-library/Tableau_Decisionnel_TEMPLATE.xlsx"
    tableau_report_path = f"/local-files/TABLEAU_DECISIONNEL_{FILE_ID}.xlsx"
    import tempfile as _tf3, json as _json3
    with _tf3.NamedTemporaryFile(suffix=".json", delete=False, mode="w", encoding="utf-8") as _tmp3:
        _json3.dump(gonogo_data, _tmp3, ensure_ascii=False)
        _tmp3_name = _tmp3.name
    # Session 36: pass the FULL batch file list (not just ao_pdf_path, the
    # first file) so tableau_decisionnel.py's bordereau detection can scan
    # every file in the batch, not just the main AO document.
    with _tf3.NamedTemporaryFile(suffix=".json", delete=False, mode="w", encoding="utf-8") as _tmp4:
        _json3.dump(ao_pdf_paths, _tmp4, ensure_ascii=False)
        _tmp4_name = _tmp4.name
    _tab_result = subprocess.run(
        ["python3", "/root/tableau_decisionnel.py", _tableau_template, _tmp3_name, _tmp4_name, tableau_report_path],
        capture_output=True, text=True, timeout=120
    )
    os.unlink(_tmp3_name)
    os.unlink(_tmp4_name)
    if _tab_result.returncode != 0:
        raise RuntimeError(f"tableau_decisionnel.py exited {_tab_result.returncode}: {_tab_result.stderr[-500:]}")
    print(f"✅ Tableau décisionnel: {tableau_report_path}")
except Exception as _et:
    print(f"⚠️ Tableau décisionnel error (non-blocking): {_et}")
    tableau_report_path = None


# ── Derive safe_name from FILE_NAME (same logic as pipeline.py) ──────────────
import re as _re
safe_name = FILE_NAME
for ext in ['.pdf', '.docx', '.doc']:
    safe_name = safe_name.replace(ext, '').replace(ext.upper(), '')
safe_name = safe_name.replace('Proposition_CHG_', '')
safe_name = _re.sub(r'[^a-zA-Z0-9_\-]', '_', safe_name)[:60]

# ── Agent 2 — Planning document ──────────────────────────────────────────────
planning_report_path = None
try:
    import tempfile as _tf2, json as _json2
    planning_report_path = f"/local-files/PLANIFICATION_{FILE_ID}.docx"
    with _tf2.NamedTemporaryFile(suffix=".json", delete=False, mode="w", encoding="utf-8") as _tmp2:
        _json2.dump(gonogo_data, _tmp2, ensure_ascii=False)
        _tmp2_name = _tmp2.name
    run_planning_agent(_tmp2_name, planning_report_path)
    os.unlink(_tmp2_name)
    print(f"\u2705 Planning document: {planning_report_path}")
except Exception as _ep:
    print(f"\u26a0\ufe0f Planning agent error: {_ep}")
    planning_report_path = None

# ── N8N_OUTPUT ────────────────────────────────────────────────────────────────
# Session 51: status/stages/degraded_reasons computed from the three existing
# *_report_path is None checks (already the de facto per-stage failure signal,
# just never surfaced into "status" before -- it was hardcoded "success"
# regardless of what happened above). "success" stays the literal string for
# full success, unchanged from today; "partial"/"error" are newly introduced,
# strictly on paths that were previously always mislabeled "success". Confirmed
# via full n8n workflow export (Session 51) that no workflow branches on this
# inner status field -- only on pipeline_server.py's own outer STATUS_*.json
# status ("done"/"error"/"processing"), which is untouched by this change.
stages = {
    "gonogo": "success" if gonogo_report_path else "error",
    "tableau": "success" if tableau_report_path else "error",
    "planning": "success" if planning_report_path else "error",
}
degraded_reasons = [f"{stage}_failed" for stage, outcome in stages.items() if outcome == "error"]
_failed_extraction_docs = gonogo_data.get("documents_extraction_failed", [])
if _failed_extraction_docs:
    degraded_reasons.append("additional_document_extraction_failed")
if stages["gonogo"] == "error":
    overall_status = "error"  # core deliverable (the Go/No-Go decision) missing
elif degraded_reasons:
    overall_status = "partial"
else:
    overall_status = "success"

result = {
    "status": overall_status,
    "stages": stages,
    "degraded_reasons": degraded_reasons,
    "documents_extraction_failed": _failed_extraction_docs,
    "ao_client": ao_data.get("client", ao_data.get("municipality", ao_data.get("municipalite", ""))),
    "ao_titre": ao_data.get("titre_projet", ao_data.get("titre", ao_data.get("title", ""))),
    "numero_ao": ao_data.get("numero_ao", ""),  # Session 39: needed for addenda lookup by exact AO number
    "gonogo_score": gonogo_score,
    "gonogo_recommandation": gonogo_recommandation,
    "gonogo_nb_incongruites": gonogo_nb_incongruites,
    "gonogo_report_file": gonogo_report_path.replace("/local-files/", "/files/") if gonogo_report_path else "",
    "planning_report_file": planning_report_path.replace("/local-files/", "/files/") if planning_report_path else "",
    "tableau_report_file": tableau_report_path.replace("/local-files/", "/files/") if tableau_report_path else "",
}
print(f"N8N_OUTPUT:{json.dumps(result)}")
